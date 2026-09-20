"""Deterministic staff-deployment workbook importer (spec docs/superpowers/specs/2026-09-19-curriculum-plan-design.md §4).

Reads a department-per-sheet workbook (one row per teaching requirement), a `Groups` sheet
(divisions and bands) and a `Control`/`Load` sheet (staff), into a curriculum plan document.
No LLM: everything here is deterministic and reports issues instead of raising on bad data.
"""
from __future__ import annotations

import copy
import csv
import hashlib
import io
import re
from datetime import datetime, timezone

from .. import names as names_mod
from . import model as M
from .issues import Issue

#  README is skipped on both shapes of workbook: the start wizard's `wizard.instantiate.workbook_bytes`
# always appends one (spec docs/superpowers/specs/2026-09-20-start-wizard-design.md §5), and it is
# never a department, staff or duties sheet.
_SKIP_SHEETS = {"master", "upload", "overall", "level of prep setup", "cca", "committees", "readme"}
_SPECIAL_SHEETS = {"groups", "control", "load"}

_NEEDED_HEADERS = {"subject", "single", "double", "grouping", "class 1"}
_GENERIC_SHEETS = {"staff", "venues", "units", "duties"}

_LESSON_COLS = (("single", "1"), ("double", "2"), ("triple", "3"), ("quadruple", "4"))

_PERIODS_SUFFIX = re.compile(r"\s*\((\d+)\s*Periods?\)\s*$", re.IGNORECASE)

_STAFF_NAME_KEYS = ("staff", "staff name")
_STAFF_SIGNAL_KEYS = ("dept", "teaching load factor", "teaching load", "short")


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------

def _cell(row, col):
    if col is None or col >= len(row):
        return None
    return row[col]


def _text(row, col) -> str:
    v = _cell(row, col)
    return str(v).strip() if v is not None else ""


def _header_map(row) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, c in enumerate(row or ()):
        if c is not None and str(c).strip():
            out[str(c).strip().lower()] = i
    return out


def _row_hash(row) -> str:
    h = hashlib.sha1()
    h.update("\x1f".join("" if v is None else str(v) for v in row).encode("utf-8"))
    return h.hexdigest()


def _family(code: str) -> str:
    """The letters following the leading digit(s) of a group code, minus any trailing digits
    and its trailing variant character: "1ELP" -> "EL", "1EMP" -> "EM", "1HCR2" -> "HC"."""
    rest = re.sub(r"^\d+", "", code)
    rest = re.sub(r"\d+$", "", rest)
    return rest[:-1] if len(rest) > 1 else rest


def _resolve_staff(name: str, org: dict | None, staff: list[dict]) -> str | None:
    """Match `name` to an existing person (`names.find`), else to an already-known staff
    entry by name, else create a new staff entry with a slug id and no availability of its own
    (`avail: None`), which generation reads as the whole cycle: only the Control sheet says when
    someone works, and a made-up window would be rejected by every grid it does not fit."""
    name = (name or "").strip()
    if not name:
        return None
    if org is not None:
        matches = names_mod.find(org, name, kinds=("person",))
        if matches:
            pid = matches[0]["id"]
            if not any(s["id"] == pid for s in staff):
                staff.append({"id": pid, "name": matches[0]["name"], "short": "", "dept": "",
                              "load_factor": 1.0, "avail": None, "max_periods_day": None, "source_hash": None})
            return pid
    low = name.lower()
    new_id = M.slug(name)
    for s in staff:
        if s["name"].strip().lower() == low or s["id"] == new_id:
            return s["id"]
    staff.append({"id": new_id, "name": name, "short": "", "dept": "", "load_factor": 1.0,
                  "avail": None, "max_periods_day": None, "source_hash": None})
    return new_id


# ---------------------------------------------------------------------------
# is_deployment_workbook
# ---------------------------------------------------------------------------

def is_deployment_workbook(data: bytes) -> bool:
    try:
        wb = openpyxl_load(data)
    except Exception:
        return False
    try:
        for ws in wb.worksheets:
            row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if row is None:
                continue
            cells = {str(c).strip().lower() for c in row if c is not None}
            if _NEEDED_HEADERS <= cells:
                return True
        return False
    finally:
        wb.close()


def openpyxl_load(data: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def is_generic_workbook(data: bytes) -> bool:
    """A duties workbook (spec §5): Staff, Venues, Units and Duties sheets, by name — the same
    shape `wizard.instantiate.workbook_bytes` writes for every non-education template."""
    try:
        wb = openpyxl_load(data)
    except Exception:
        return False
    try:
        titles = {ws.title.strip().lower() for ws in wb.worksheets}
        return _GENERIC_SHEETS <= titles
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# staff (Control / Load sheet)
# ---------------------------------------------------------------------------

def _read_staff_sheet(ws, org: dict | None) -> list[dict]:
    staff: list[dict] = []
    rows = list(ws.iter_rows(values_only=True))
    header = None
    header_idx = None
    for i, row in enumerate(rows):
        hm = _header_map(row)
        if any(k in hm for k in _STAFF_NAME_KEYS) and any(k in hm for k in _STAFF_SIGNAL_KEYS):
            header, header_idx = hm, i
            break
    if header is None:
        return staff

    name_col = header.get("staff", header.get("staff name"))
    dept_col = header.get("dept")
    load_col = header.get("teaching load factor", header.get("teaching load"))
    short_col = header.get("short")

    for row in rows[header_idx + 1:]:
        name = _text(row, name_col)
        if not name:
            continue
        pid = None
        if org is not None:
            matches = names_mod.find(org, name, kinds=("person",))
            if matches:
                pid = matches[0]["id"]
        if pid is None:
            pid = M.slug(name)
        load_raw = _cell(row, load_col)
        staff.append({
            "id": pid,
            "name": name,
            "short": _text(row, short_col),
            "dept": _text(row, dept_col),
            "load_factor": float(load_raw) if load_raw not in (None, "") else 1.0,
            "avail": None,
            "max_periods_day": None,
            "source_hash": _row_hash(row),
        })
    return staff


# ---------------------------------------------------------------------------
# department sheets (requirements)
# ---------------------------------------------------------------------------

_SPLIT = re.compile(r"^\s*(\d+)\s*/\s*(\d+)\s*$")


def _apply_split(req: dict, split_raw: str, teachers: list[str], periods, issues: list[Issue]) -> list[dict]:
    """A `Split` cell ("4/2"): the row's two teachers no longer co-teach every lesson, each takes
    their own share of the periods instead — one requirement a teacher, carried as a single lesson
    of that many slots (plan.model.LESSON_LENGTH_MAX allows it). The two numbers must sum to the
    row's periods; when they do not, the row is left as the single co-taught requirement it already
    is and a block issue is reported instead of guessing how to split it."""
    m = _SPLIT.match(split_raw)
    a, b = (int(m.group(1)), int(m.group(2))) if m else (None, None)
    if a is None or periods is None or a + b != periods:
        issues.append(Issue("block", req["id"],
                             f"{req['id']}: split {split_raw!r} does not sum to periods {periods}"))
        return [req]
    return [
        {**req, "id": f"{req['id']}-a", "teachers": [teachers[0]], "lessons": {str(a): 1}, "periods": a},
        {**req, "id": f"{req['id']}-b", "teachers": [teachers[1]], "lessons": {str(b): 1}, "periods": b},
    ]


def _read_department_sheet(ws, org, staff, issues: list[Issue]) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = _header_map(rows[0])
    if "subject" not in header:
        issues.append(Issue("warn", ws.title, f"{ws.title}: no 'Subject' column found, sheet skipped"))
        return []

    dept = ws.title
    subject_col = header["subject"]
    level_col = header.get("subject level", header.get("level"))
    total_col = header.get("total")
    grouping_col = header.get("grouping")
    size_col = header.get("total students")
    class_cols = [header.get(f"class {n}") for n in range(1, 7)]
    teacher_cols = sorted(
        (int(k.split(" ", 1)[1]), v) for k, v in header.items()
        if k.startswith("teacher ") and k.split(" ", 1)[1].strip().isdigit()
    )
    split_col = header.get("split")

    requirements: list[dict] = []
    for row in rows[1:]:
        subject_raw = _text(row, subject_col)
        if not subject_raw:
            continue

        m = _PERIODS_SUFFIX.search(subject_raw)
        if m:
            subject = subject_raw[:m.start()].strip()
            suffix_periods = int(m.group(1))
        else:
            subject = subject_raw
            suffix_periods = None

        total_val = _cell(row, total_col)
        if total_val not in (None, ""):
            periods = int(total_val)
        elif suffix_periods is not None:
            periods = suffix_periods
        else:
            periods = None

        lessons = {}
        for name_key, lkey in _LESSON_COLS:
            v = _cell(row, header.get(name_key))
            lessons[lkey] = int(float(v)) if v not in (None, "") else 0

        classes = [_text(row, c) for c in class_cols if c is not None and _text(row, c)]
        level = _text(row, level_col)
        grouping_raw = _text(row, grouping_col)
        is_entire_class = grouping_raw.strip().lower() == "entire class"

        teachers: list[str] = []
        for _, col in teacher_cols:
            tname = _text(row, col)
            if not tname:
                continue
            tid = _resolve_staff(tname, org, staff)
            if tid and tid not in teachers:
                teachers.append(tid)

        size_val = _cell(row, size_col)
        size = int(size_val) if size_val not in (None, "") else None
        split_raw = _text(row, split_col)

        row_hash = _row_hash(row)

        def make_req(req_classes: list[str], grouping: str) -> dict:
            return {
                "id": M.requirement_id(dept, level, grouping, req_classes),
                "dept": dept,
                "level": level,
                "subject": subject,
                "periods": periods,
                "lessons": dict(lessons),
                "grouping": grouping,
                "classes": req_classes,
                "teachers": list(teachers),
                "size": size,
                "venue": {"kind": None, "room": None},
                "sync_with": None,
                "source_hash": row_hash,
            }

        if not classes:
            continue

        if is_entire_class:
            row_reqs = [make_req([cls], "class") for cls in classes]
        else:
            row_reqs = [make_req(classes, grouping_raw)]

        if split_raw and len(row_reqs) == 1 and len(teachers) >= 2:
            row_reqs = _apply_split(row_reqs[0], split_raw, teachers[:2], periods, issues)
        requirements.extend(row_reqs)

        if periods is not None:
            expected = sum(int(k) * v for k, v in lessons.items())
            if periods != expected:
                for req in row_reqs:
                    issues.append(Issue("block", req["id"],
                                         f"{req['id']}: periods {periods} does not match lessons total {expected}"))

    return requirements


# ---------------------------------------------------------------------------
# Groups sheet (divisions and bands)
# ---------------------------------------------------------------------------

_GROUP_NUMBER = re.compile(r"^group\s*\d+$", re.IGNORECASE)
_PLAIN_NUMBER = re.compile(r"^\d+(\.0*)?$")


def _is_group_subheader(row, class_cols, group_cols) -> bool:
    """A second header row of the Groups sheet: it numbers the group columns ("Group 1", "Group 2")
    or the class columns ("1", "2", …), where a division row names classes ("1A1") and group
    codes ("1ELP")."""
    codes = [_text(row, c) for c in group_cols if _text(row, c)]
    if codes and all(_GROUP_NUMBER.match(t) for t in codes):
        return True
    classes = [_text(row, c) for c in class_cols if _text(row, c)]
    return bool(classes) and all(_PLAIN_NUMBER.match(t) for t in classes)


def _read_groups_sheet(ws, requirements: list[dict]) -> tuple[list[dict], list[dict], list[Issue]]:
    divisions: list[dict] = []
    bands: list[dict] = []
    issues: list[Issue] = []

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return divisions, bands, issues

    header = rows[0] or ()
    class_start = groups_start = None
    for i, c in enumerate(header):
        if c is None:
            continue
        text = str(c).strip().lower()
        if text.startswith("class") and class_start is None:
            class_start = i
        elif text.startswith("groups") and groups_start is None:
            groups_start = i
    if class_start is None or groups_start is None:
        issues.append(Issue("warn", ws.title, f"{ws.title}: could not find the division/group header row"))
        return divisions, bands, issues

    class_cols = range(class_start, groups_start)
    max_len = max((len(r) for r in rows), default=groups_start)
    group_cols = range(groups_start, max_len)

    by_grouping: dict[str, list[str]] = {}
    for r in requirements:
        by_grouping.setdefault(r["grouping"], []).append(r["id"])

    # The divisions start after the header rows, however many there are: the school's sheet numbers
    # its class columns ("1", "2", …) and its group columns ("Group 1", "Group 2") on a second row,
    # but a sheet with one header row must not lose the division on its first data row.
    start = 1
    while start < len(rows) and _is_group_subheader(rows[start], class_cols, group_cols):
        start += 1

    for row in rows[start:]:
        classes = [_text(row, c) for c in class_cols if _text(row, c)]
        if not classes:
            continue
        codes = [_text(row, c) for c in group_cols if _text(row, c)]

        div_id = M.division_id(classes)
        divisions.append({"id": div_id, "classes": classes})

        families: dict[str, list[str]] = {}
        for code in codes:
            families.setdefault(_family(code), []).append(code)

        for fam, fam_codes in families.items():
            options: list[str] = []
            for code in fam_codes:
                matched = by_grouping.get(code, [])
                if not matched:
                    issues.append(Issue("warn", code, f"group code {code!r} matches no requirement"))
                options.extend(matched)
            if options:
                bands.append({"id": M.band_id(div_id, fam), "division": div_id, "options": options})

    return divisions, bands, issues


# ---------------------------------------------------------------------------
# CSV sizes
# ---------------------------------------------------------------------------

def read_sizes_csv(text: str) -> dict[str, int]:
    out: dict[str, int] = {}
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return out
    header = [c.strip().lower() for c in rows[0]]

    def idx(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    group_col = idx("teaching group")
    class_col = idx("class")
    enrolled_col = idx("enrolled")
    capacity_col = idx("capacity")

    def cell(row, i):
        return row[i].strip() if i is not None and i < len(row) and row[i] is not None else ""

    for row in rows[1:]:
        if not row or not any(c and c.strip() for c in row):
            continue
        value = cell(row, enrolled_col) or cell(row, capacity_col)
        if not value:
            continue
        try:
            size = int(float(value))
        except ValueError:
            continue
        # a CSV may carry both columns (one row naming a teaching group and the class it draws from):
        # take every key the row offers rather than only the first column that exists
        if group_col is not None:
            code = cell(row, group_col)
            key = code.split("-")[-1].strip() if code else ""
            if key:
                out[key] = size
        if class_col is not None:
            code = cell(row, class_col)
            if code:
                out[code] = size
    return out


def apply_sizes(plan: dict, sizes: dict[str, int]) -> dict:
    plan = copy.deepcopy(plan)
    for r in plan.get("requirements", []):
        key = r.get("grouping")
        if key and key != "class" and key in sizes:
            r["size"] = sizes[key]
    for c in plan.get("classes", []):
        code = c.get("code")
        if code in sizes:
            c["size"] = sizes[code]
    return M.normalise(plan)


# ---------------------------------------------------------------------------
# read_workbook
# ---------------------------------------------------------------------------

def read_workbook(data: bytes, org: dict | None, existing: dict | None, filename: str = "") -> tuple[dict, list[Issue]]:
    issues: list[Issue] = []
    try:
        wb = openpyxl_load(data)
    except Exception as e:
        issues.append(Issue("block", "workbook", f"could not read workbook: {e}"))
        base = existing if existing is not None else M.empty_plan()
        return M.normalise(base), issues

    try:
        control_ws = next((ws for ws in wb.worksheets if ws.title.strip().lower() == "control"), None)
        if control_ws is None:
            control_ws = next((ws for ws in wb.worksheets if ws.title.strip().lower() == "load"), None)
        staff = _read_staff_sheet(control_ws, org) if control_ws is not None else []

        requirements: list[dict] = []
        for ws in wb.worksheets:
            title_l = ws.title.strip().lower()
            if title_l in _SKIP_SHEETS or title_l.startswith("sheet") or title_l in _SPECIAL_SHEETS:
                continue
            requirements.extend(_read_department_sheet(ws, org, staff, issues))

        groups_ws = next((ws for ws in wb.worksheets if ws.title.strip().lower() == "groups"), None)
        if groups_ws is not None:
            divisions, bands, group_issues = _read_groups_sheet(groups_ws, requirements)
            issues.extend(group_issues)
        else:
            divisions, bands = [], []

        class_codes = sorted({c for r in requirements for c in r["classes"]})
        classes = [{"code": c, "level": "", "size": None} for c in class_codes]

        plan = {
            "version": M.PLAN_VERSION,
            "staff": staff,
            "classes": classes,
            "divisions": divisions,
            "requirements": requirements,
            "bands": bands,
            "rules": {"edge_subjects": [], "no_double_across_rest": True, "pinned": []},
            "source": {"file": filename, "imported": datetime.now(timezone.utc).isoformat()},
        }
        new_plan = M.normalise(plan)
    finally:
        wb.close()

    if existing is not None:
        return M.merge_import(existing, new_plan), issues
    return new_plan, issues


# ---------------------------------------------------------------------------
# read_generic_workbook (spec §5): a duties workbook — Staff, Venues, Units, Duties, README
# ---------------------------------------------------------------------------

_DAY_NAMES = ("monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday")
_DAY_WORDS = {"weekday": (0, 1, 2, 3, 4), "weekdays": (0, 1, 2, 3, 4),
              "weekend": (5, 6), "weekends": (5, 6)}
# Whole-cycle words: what a person writes when they are simply always available.
_WHOLE_CYCLE = {"all", "all week", "all days", "any", "anytime", "always", "every day", "everyday",
                "daily", "full", "full time", "full-time", "whole cycle", "-", "n/a", "na"}
_SPLIT_PARTS = re.compile(r"[;,&/]|\band\b|\+", re.IGNORECASE)
_RANGE = re.compile(r"\s*(?:-|–|—|\bto\b)\s*", re.IGNORECASE)


def _day_index(token: str) -> int | None:
    """"Mon", "mon.", "tues", "Wednesday" -> 0, 1, 2. Three letters is the shortest that is a day."""
    t = token.strip().strip(".").lower()
    if len(t) < 3:
        return None
    return next((i for i, name in enumerate(_DAY_NAMES) if name.startswith(t)), None)


def _days_of(part: str) -> tuple[int, ...] | None:
    """The weekdays one comma-free piece names: "Mon" -> (0,), "Mon-Wed" -> (0, 1, 2), "weekdays" ->
    (0..4). `None` when the piece is not weekdays at all."""
    word = part.strip().strip(".").lower()
    if word in _DAY_WORDS:
        return _DAY_WORDS[word]
    ends = _RANGE.split(part, maxsplit=1)
    if len(ends) == 2:
        first, last = _day_index(ends[0]), _day_index(ends[1])
        if first is None or last is None:
            return None
        span = range(first, last + 1) if first <= last else [*range(first, 7), *range(0, last + 1)]
        return tuple(span)
    one = _day_index(part)
    return None if one is None else (one,)


def _numeric_window(part: str) -> list[int] | None:
    ends = _RANGE.split(part.strip(), maxsplit=1)
    if len(ends) != 2:
        return None
    try:
        start, end = int(ends[0].strip()), int(ends[1].strip())
    except ValueError:
        return None
    return [start, end] if end > start else None


def _merge(windows: list[list[int]], n_slots: int | None) -> list[list[int]]:
    """Sorted, non-overlapping and inside the cycle — the shape `model.normalise_avail` accepts."""
    clipped = []
    for start, end in windows:
        start, end = max(0, start), end if n_slots is None else min(end, n_slots)
        if end > start:
            clipped.append([start, end])
    out: list[list[int]] = []
    for start, end in sorted(clipped):
        if out and start <= out[-1][1]:
            out[-1][1] = max(out[-1][1], end)
        else:
            out.append([start, end])
    return out


def _parse_availability(text: str, cycle: tuple[int, int, int] | None = None,
                        where: str = "", issues: list[Issue] | None = None) -> list[list[int]] | None:
    """When someone works, in the two forms the workbook's own README offers (spec §5):

    * whole days by name — "Mon-Wed", "Mon, Wed, Fri", "weekdays", and "All week"/"" for the whole
      cycle — mapped onto the cycle's days through `cycle` (slots a day, days in the cycle, days in
      a week): a 7-day-week rota reads Mon..Sun, a 5-day cycle Mon..Fri, and a multi-week cycle
      applies the days to every week;
    * slot windows in the app's own words, "0-24; 48-72", for anyone who wants the slots themselves.

    Anything else is a warn issue naming the row, and reads as the whole cycle: an unreadable cell
    should say so and carry on, not block the import."""
    text = (text or "").strip()
    if not text or text.lower() in _WHOLE_CYCLE:
        return None

    def unreadable() -> None:
        if issues is not None:
            issues.append(Issue("warn", where or "staff",
                                 f"{where or 'staff'}: availability {text!r} is not a list of days "
                                 f'("Mon-Wed", "Mon, Wed, Fri", "All week") or slot windows '
                                 f'("0-24; 48-72"); read as the whole cycle'))

    spd, n_days, per_week = cycle or (0, 0, 0)
    windows: list[list[int]] = []
    for part in _SPLIT_PARTS.split(text):
        part = part.strip()
        if not part:
            continue
        window = _numeric_window(part)
        if window is not None:
            windows.append(window)
            continue
        days = _days_of(part)
        if days is None or not spd:
            unreadable()
            return None
        for week in range(-(-n_days // per_week) if per_week else 1):
            for day in days:
                index = week * per_week + day
                if index < n_days:
                    windows.append([index * spd, (index + 1) * spd])
    merged = _merge(windows, spd * n_days if spd and n_days else None)
    if not merged:
        unreadable()
        return None
    return merged


def cycle_shape(org: dict | None, settings: dict | None) -> tuple[int, int, int] | None:
    """(slots a day, days in the cycle, days in a week) for the Availability column, from the
    timetable's settings or, failing that, the live organisation.

    The week comes from `time.days_per_week`, the same number that names the days ("Odd Mon", "Even
    Tue"), so a day named here is the day the timetable shows. Only a settings document old enough
    not to carry it is guessed at, by the length of the cycle: whole weeks run Mon..Sun, a cycle
    that divides by five Mon..Fri, anything else is one week of its own length."""
    time = (settings or {}).get("time") or {}
    spd = time.get("slots_per_day")
    n_slots = len(time.get("labels") or ())
    per_week = time.get("days_per_week")
    if not spd and org:
        spd = ((org.get("rules") or {}).get("slots_per_day"))
        n_slots = len(org.get("time_labels") or ())
        per_week = None                 # an organisation does not carry the week
    if not spd or not n_slots:
        return None
    spd = int(spd)
    n_days = max(1, n_slots // spd)
    if not isinstance(per_week, int) or isinstance(per_week, bool) or not 1 <= per_week <= 7:
        per_week = 7 if n_days % 7 == 0 else (5 if n_days % 5 == 0 else n_days)
    return spd, n_days, per_week


def _co_staffed_count(raw: str) -> int:
    """"Co-staffed": "" or "no" -> one person on the duty alone; "yes" -> two; a number -> that many."""
    text = (raw or "").strip().lower()
    if not text or text == "no":
        return 1
    if text == "yes":
        return 2
    try:
        n = int(float(text))
    except ValueError:
        return 1
    return n if n > 0 else 1


def _read_generic_staff(ws, org: dict | None, cycle: tuple[int, int, int] | None,
                        issues: list[Issue]) -> list[dict]:
    staff: list[dict] = []
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return staff
    header = _header_map(rows[0])
    name_col = header.get("name")
    role_col = header.get("role")
    avail_col = header.get("availability")
    load_col = header.get("load factor")

    for row in rows[1:]:
        name = _text(row, name_col)
        if not name:
            continue
        pid = None
        if org is not None:
            matches = names_mod.find(org, name, kinds=("person",))
            if matches:
                pid = matches[0]["id"]
        if pid is None:
            pid = M.slug(name)
        load_raw = _cell(row, load_col)
        staff.append({
            "id": pid,
            "name": name,
            "short": "",
            # `dept` doubles as the row's "Role" here: `_role_index` groups staff by it so
            # "Who can do it" can name a role instead of listing everyone in it.
            "dept": _text(row, role_col),
            "load_factor": float(load_raw) if load_raw not in (None, "") else 1.0,
            "avail": _parse_availability(_text(row, avail_col), cycle, name, issues),
            "max_periods_day": None,
            "source_hash": _row_hash(row),
        })
    return staff


def _role_index(staff: list[dict]) -> dict[str, list[str]]:
    index: dict[str, list[str]] = {}
    for s in staff:
        role = (s.get("dept") or "").strip().lower()
        if role:
            index.setdefault(role, []).append(s["id"])
    return index


def _read_generic_units(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = _header_map(rows[0])
    name_col = header.get("name")
    size_col = header.get("size")

    units = []
    for row in rows[1:]:
        name = _text(row, name_col)
        if not name:
            continue
        size_val = _cell(row, size_col)
        units.append({"code": name, "level": "", "size": int(size_val) if size_val not in (None, "") else None})
    return units


def _who_can_do_it(who_raw: str, org, staff: list[dict], role_index: dict[str, list[str]]) -> list[str]:
    """A comma list of staff names, or a role from the Staff sheet's "Role" column: a role expands
    to everyone with it, in Staff-sheet order; a name resolves the same way a deployment workbook's
    teacher columns do (`_resolve_staff`), so an unlisted name still becomes a staff entry."""
    pool: list[str] = []
    for token in (who_raw or "").split(","):
        token = token.strip()
        if not token:
            continue
        role_ids = role_index.get(token.lower())
        if role_ids:
            for pid in role_ids:
                if pid not in pool:
                    pool.append(pid)
            continue
        pid = _resolve_staff(token, org, staff)
        if pid and pid not in pool:
            pool.append(pid)
    return pool


def _dedup_id(req_id: str, seen: dict[str, int]) -> str:
    """Task 2's own sample workbook cycles a short duty list across three example rows, so the same
    (Name, Unit) pair can appear twice; a repeat gets a "-2", "-3"… suffix rather than silently
    replacing the requirement that came before it."""
    seen[req_id] = seen.get(req_id, 0) + 1
    n = seen[req_id]
    return req_id if n == 1 else f"{req_id}-{n}"


def _parse_duty_count(raw, label: str, duty_name: str, req_id: str, issues: list[Issue]) -> int | None:
    """A Duties count cell ("Per cycle", "Length"): a blank cell is 0 (no lessons), a non-numeric
    one is a block issue rather than a crash — same "report, don't raise" convention as
    `_co_staffed_count` — and the caller carries `None` on, which reads as "no lessons" downstream
    (`periods`/`lessons` both fall back to empty) rather than guessing a count."""
    if raw in (None, ""):
        return 0
    try:
        return int(raw)
    except (TypeError, ValueError):
        issues.append(Issue("block", req_id, f"{duty_name}: {label} must be a whole number"))
        return None


def _read_generic_duties(ws, org, staff: list[dict], role_index: dict[str, list[str]],
                         issues: list[Issue]) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = _header_map(rows[0])
    name_col = header.get("name")
    unit_col = header.get("unit")
    per_cycle_col = header.get("per cycle")
    length_col = header.get("length")
    venue_col = header.get("venue kind")
    who_col = header.get("who can do it")
    together_col = header.get("together with")
    co_col = header.get("co-staffed")

    requirements: list[dict] = []
    by_key: dict[tuple[str, str], dict] = {}
    pending_sync: list[tuple[dict, str, str]] = []
    seen_ids: dict[str, int] = {}

    for row in rows[1:]:
        name = _text(row, name_col)
        if not name:
            continue
        unit = _text(row, unit_col)
        venue_kind = _text(row, venue_col) or None
        who_raw = _text(row, who_col)
        together = _text(row, together_col)

        req_id = _dedup_id(M.requirement_id(name, None, "class", [unit] if unit else []), seen_ids)

        per_cycle = _parse_duty_count(_cell(row, per_cycle_col), "Per cycle", name, req_id, issues)
        if per_cycle is not None and per_cycle <= 0:
            issues.append(Issue("block", req_id, f"{name}: Per cycle must be greater than zero"))
        length = _parse_duty_count(_cell(row, length_col), "Length", name, req_id, issues)

        pool = _who_can_do_it(who_raw, org, staff, role_index)
        if not pool:
            issues.append(Issue("block", req_id, f"no one can do {name}"))
        n = _co_staffed_count(_text(row, co_col))
        if pool and len(pool) < n:
            issues.append(Issue("warn", req_id, f"{name}: co-staffed {n} but only {len(pool)} eligible"))
        teachers = pool[:n]

        req = {
            "id": req_id,
            "dept": name,
            "level": "",
            "subject": name,
            "periods": per_cycle * length if per_cycle and length else None,
            "lessons": {str(length): per_cycle} if length and per_cycle else {},
            "grouping": "class",
            "classes": [unit] if unit else [],
            "teachers": teachers,
            "size": None,
            "venue": {"kind": venue_kind, "room": None},
            "sync_with": None,
            "source_hash": _row_hash(row),
        }
        requirements.append(req)
        by_key[(unit, name)] = req
        if together:
            pending_sync.append((req, unit, together))

    for req, unit, together_name in pending_sync:
        target = by_key.get((unit, together_name))
        if target is not None:
            req["sync_with"] = target["id"]
        else:
            issues.append(Issue("warn", req["id"],
                                 f"{req['id']}: 'together with' names an unknown duty {together_name!r}"))

    return requirements


def read_generic_workbook(data: bytes, org: dict | None, existing: dict | None,
                          filename: str = "", settings: dict | None = None) -> tuple[dict, list[Issue]]:
    """The generic duties workbook (spec §5): Staff, Venues, Units and Duties sheets, detected by
    name (`is_generic_workbook`) rather than by header cells. `Venues` only needs to exist for the
    shape to be recognised — a duty names its venue kind directly, the same field every requirement
    already carries."""
    issues: list[Issue] = []
    try:
        wb = openpyxl_load(data)
    except Exception as e:
        issues.append(Issue("block", "workbook", f"could not read workbook: {e}"))
        base = existing if existing is not None else M.empty_plan()
        return M.normalise(base), issues

    try:
        sheets = {ws.title.strip().lower(): ws for ws in wb.worksheets}
        cycle = cycle_shape(org, settings)
        staff = _read_generic_staff(sheets["staff"], org, cycle, issues) if "staff" in sheets else []
        role_index = _role_index(staff)
        units = _read_generic_units(sheets["units"]) if "units" in sheets else []
        requirements = (_read_generic_duties(sheets["duties"], org, staff, role_index, issues)
                        if "duties" in sheets else [])

        unit_sizes = {u["code"]: u["size"] for u in units}
        unit_codes = sorted({u["code"] for u in units} | {c for r in requirements for c in r["classes"]})
        classes = [{"code": code, "level": "", "size": unit_sizes.get(code)} for code in unit_codes]

        plan = {
            "version": M.PLAN_VERSION,
            "staff": staff,
            "classes": classes,
            "divisions": [],
            "requirements": requirements,
            "bands": [],
            "rules": {"edge_subjects": [], "no_double_across_rest": True, "pinned": []},
            "source": {"file": filename, "imported": datetime.now(timezone.utc).isoformat()},
        }
        new_plan = M.normalise(plan)
    finally:
        wb.close()

    if existing is not None:
        return M.merge_import(existing, new_plan), issues
    return new_plan, issues
