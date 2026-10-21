"""Every file a reading (spec docs/superpowers/specs/2026-10-12-assistant-files-design.md §2.1, §2.4).

`read_files` routes each file: an aSc export (all of a drop's, together) → one timetable reading; a
planning workbook or a class-sizes CSV → an exact plan reading (merged by the importer, nothing
stored; a drop's sizes CSVs go into the reading of its first planning workbook, applied on top); a
student list → a students reading, read exactly by students.read_rows (one that cannot be read
exactly is a failed card: a student list never reaches the AI); a workbook's Combinations sheet →
its own combinations reading, read exactly (never by the AI) and applied into `blocks_setup`; one of the five
CSV files (staff, venues, requirements, combinations, the level students list: csv-pack spec) → its own exact
reading, a drop's read in dependency order whatever the upload order; any other .xlsx/.csv → the smart
reader; .docx and other PDFs → the document reader; anything else → a failed card that says why. A file the model could not read now (or read as nothing) is kept
on disk for Try again. Nothing here writes the plan or the draft."""
from __future__ import annotations

import csv
import io
import logging
import re
import secrets
import time
from pathlib import Path

from ... import asc_import, blocks as BK, students as ST
from ...config import MAX_UPLOAD
from ...extract import UnsupportedFile, extract
from ...intake import IntakeError, apply_patch as apply_org_patch
from ...llm import ProviderError
from ... import week as W
from ... import levels as LV
from .. import importer as IMP
from .. import model as M
from .. import specialisation as SP
from . import csvpack as CP, document, packmerge as PM, routes as RR, sections as S, writer
from . import recipes as R
from .render import WorkbookError
from .tabular import cell as _cell, decoded as _decoded, delimiters as _delimiters, head as _head

FILES_MAX = 10
TYPES = (".pdf", ".docx", ".xlsx", ".csv")
SWEEP_GRACE = 600            # seconds a kept file is left alone by the sweep: a drop being read may not have stored it yet
OPTIONS = {"mode": "keep", "classes_as_planes": False}
WAITING = ("waiting", "failed", "retry")      # what clear drops; applied and skipped readings stay
_log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# kept files
# ---------------------------------------------------------------------------

def files_dir(db) -> Path:
    """Where readings keep their files: beside the chat's uploads, under the app's data directory."""
    d = Path(db.path).parent / "uploads" / "readings"
    d.mkdir(parents=True, exist_ok=True)
    return d


def keep_file(db, name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    path = files_dir(db) / f"{secrets.token_hex(12)}{ext if ext in TYPES else '.bin'}"
    path.write_bytes(data)
    return str(path)


def _paths(reading: dict) -> list[str]:
    return [p for p in [reading.get("path"), *(reading.get("paths") or []), *(reading.get("sizes_paths") or [])]
            if isinstance(p, str) and p]


def _unlink(path: str) -> None:
    try:
        Path(path).unlink(missing_ok=True)
    except OSError:
        pass


def _named(db) -> set[str]:
    """Every file a reading of any timetable names."""
    out: set[str] = set()
    for t in db.timetables():
        for e in RR._stored(db, t["id"]):
            out.update(_paths(e["reading"]))
    return out


def sweep(db) -> None:
    """Unlink kept files no reading names (a deleted timetable's, a replaced or trimmed reading's)."""
    named = {str(Path(p)) for p in _named(db)}
    cutoff = time.time() - SWEEP_GRACE
    for f in files_dir(db).iterdir():
        try:
            if f.is_file() and str(f) not in named and f.stat().st_mtime < cutoff:
                f.unlink()
        except OSError:
            pass


# ---------------------------------------------------------------------------
# storing a reading
# ---------------------------------------------------------------------------

def _entry(db, name: str, reading: dict, tid: str | None) -> dict:
    return RR.store(db, name, reading, tid)


def _failed(db, name: str, message: str, tid: str | None, path: str | None = None, status: str = "failed") -> dict:
    entry = _entry(db, name, {"kind": "plan", "via": "ai", "status": status, "sections": {}, "doubts": [],
                              "summary": name, "error": message, "path": path}, tid)
    return {"name": name, "status": status, "id": entry["id"], "message": message}


def _current(db, tid: str | None) -> dict:
    return M.normalise(db.get_value("plan", tid) or M.empty_plan())


def _plan_reading(db, name: str, body: dict, via: str, tid: str | None, data: bytes | None = None) -> dict:
    """A plan reading stored to wait for Apply. One read by the AI (`data`, its file) that brings nothing
    is kept to try again: an empty reply is the model's, not the file's."""
    current = _current(db, tid)
    try:
        secs = S.sections_of(current, S.merged(current, {**body, "file": name}))
    except M.PlanError as e:
        return _failed(db, name, f"What was read from {name} does not fit the plan: {e}", tid)
    if not secs and via == "ai" and data is not None:
        return _retry(db, name, data, f"Nothing usable came back when the AI read {name}: try again, "
                                      "or check that it adds something to the plan.", tid)
    if not secs:
        why = next(iter(body.get("doubts") or []), "")
        return _failed(db, name, f"Nothing in {name} adds to the plan{': ' + why if why else '.'}", tid)
    stored = {k: v for k, v in body.items() if k != "merged"}
    entry = _entry(db, name, {**stored, "kind": "plan", "via": via, "status": "waiting", "sections": secs}, tid)
    summary = body.get("summary") or name
    return {"name": name, "status": "read", "id": entry["id"],
            "message": f"{summary}: " + ", ".join(f"{s['count']} {section}" for section, s in secs.items())}


def _timetable_reading(db, name: str, body: dict, ticked: bool, tid: str | None) -> dict:
    org = body.get("organisation") or body.pop("_org")
    secs = {"timetable": {"count": len(org.get("events") or []), "teachers": len(org.get("persons") or []),
                          "samples": [], "ticked": ticked}}
    reading = {k: v for k, v in body.items() if k != "_org"}
    entry = _entry(db, name, {**reading, "kind": "timetable", "via": "ai" if body.get("organisation") else "exact",
                              "status": "waiting", "sections": secs, "options": dict(OPTIONS)}, tid)
    t = secs["timetable"]
    return {"name": name, "status": "read", "id": entry["id"],
            "message": f"{body.get('summary') or name}: {t['count']} lessons, {t['teachers']} teachers"}


# ---------------------------------------------------------------------------
# the readers
# ---------------------------------------------------------------------------

def sizes_text(name: str, data: bytes) -> str | None:
    """A class-sizes CSV's text (a Class or Teaching Group column with Enrolled, or a Teaching Group column
    with Capacity, and a size in some row); None for any other file. A Class column with Capacity alone
    is a room's capacity (a home-room list), not a class size."""
    if not name.lower().endswith(".csv"):
        return None
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    try:
        header = {c.strip().lower() for c in next(csv.reader(io.StringIO(text)), [])}
    except csv.Error:                       # NULs or a stray newline: not a sizes CSV
        return None
    if not ("enrolled" in header and header & {"class", "teaching group"}
            or {"capacity", "teaching group"} <= header):
        return None
    return text if IMP.read_sizes_csv(text) else None


def is_planning(name: str, data: bytes) -> bool:
    """A deployment or duties workbook: read exactly by the importer."""
    return name.lower().endswith(".xlsx") and (IMP.is_deployment_workbook(data) or IMP.is_generic_workbook(data))


def _with_sizes(plan: dict, texts: list[str]) -> tuple[dict, list[str]]:
    """`plan` with the sizes of `texts` (sizes CSVs) on top, and a doubt naming sizes that fit nothing."""
    sizes: dict[str, int] = {}
    for text in texts:
        sizes.update(IMP.read_sizes_csv(text))
    if not sizes:
        return plan, []
    known = {c["code"] for c in plan["classes"]} | {r.get("grouping") for r in plan["requirements"]}
    unknown = [k for k in sizes if k not in known]
    doubts = [f"{len(unknown)} size{'' if len(unknown) == 1 else 's'} name no class or group in the plan: "
              + ", ".join(unknown[:8])] if unknown else []
    return IMP.apply_sizes(plan, sizes), doubts


def exact_merged(db, name: str, data: bytes, current: dict, tid: str | None,
                 sizes: list[tuple[str, bytes]] = ()) -> tuple[dict, list[str]]:
    """A planning workbook (or a sizes CSV) merged into `current` by the importer: (plan, its issue texts).
    `sizes`, the (name, bytes) of sizes CSVs given with a workbook, are applied on top of its merge."""
    text = sizes_text(name, data)
    if text is not None:
        return _with_sizes(current, [text])
    live = db.get_org("live", tid)
    if IMP.is_generic_workbook(data):
        merged, issues = IMP.read_generic_workbook(data, live, current, name, db.get_settings(tid))
    else:
        merged, issues = IMP.read_workbook(data, live, current, name)
    texts = [t for t in (sizes_text(n, d) for n, d in sizes) if t is not None]
    merged, size_doubts = _with_sizes(merged, texts)
    return merged, list(dict.fromkeys([*(i.text for i in issues), *size_doubts]))[:document.DOUBTS_MAX]


NOT_KEPT = "the file wasn't kept because its Students or Combinations sheet couldn't be removed"


def _kept_workbook(name: str, data: bytes) -> bytes | None:
    return without_students_sheet(data) if name.lower().endswith(".xlsx") else data


def _open_for_copy(data: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)    # values: a saved formula would lose its cached result


def _is_kept_out(title) -> bool:
    """Whether a sheet title is one of the sheets read once at the drop and never kept (Students, Combinations)."""
    return " ".join(str(title).split()).casefold() in KEPT_OUT


def _names_students_sheet(data: bytes) -> bool:
    """Whether the workbook's own list of sheets (xl/workbook.xml) names a Students or Combinations sheet:
    asked without openpyxl, for a workbook it could not open."""
    import html
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            book = z.read("xl/workbook.xml").decode("utf-8", "replace")
    except Exception:
        return False
    return any(_is_kept_out(html.unescape(n))
               for n in re.findall(r"<(?:\w+:)?sheet\b[^>]*?\bname=\"([^\"]*)\"", book))


def _has_other_sheets(data: bytes) -> bool:
    """Whether the workbook has a sheet besides Students and Combinations. True when it cannot be told (the
    stripping then decides, failing closed)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    except Exception:
        return True
    try:
        return any(not _is_kept_out(ws.title) for ws in wb.worksheets)
    finally:
        wb.close()


def without_students_sheet(data: bytes) -> bytes | None:
    """`data` (a planning workbook) without any sheet named Students or Combinations: the copy that is kept
    for Apply. The students are read once, at the drop, and live only in the timetable's `students` value
    (and a reading until it is applied or skipped); the combinations likewise end in `blocks_setup`; the
    plan reading never needs either, so they are not left on disk in a file nothing removes. A workbook
    with no such sheet is kept as it came. Fails closed: None (keep no copy) when the workbook has such a
    sheet that could not be taken out."""
    try:
        wb = _open_for_copy(data)
    except Exception:
        return None if _names_students_sheet(data) else data
    try:
        extra = [ws for ws in wb.worksheets if _is_kept_out(ws.title)]
        if not extra:
            return data
        for ws in extra:
            wb.remove(ws)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception:
        return None
    finally:
        wb.close()


def _exact(db, name: str, data: bytes, tid: str | None, summary: str = "Planning workbook",
           sizes: list[tuple[str, bytes]] = (), alongside: list | None = None) -> dict:
    """A planning workbook (or sizes CSV) as an exact plan reading. When `alongside` is a list and the
    workbook's Students sheet holds students, their reading (read exactly, against the plan this workbook
    brings) is appended to it: the one file then makes two readings, a plan and a student list; its
    Combinations sheet, when it holds any, likewise makes a combinations reading."""
    merged, doubts = exact_merged(db, name, data, _current(db, tid), tid, sizes)
    kept = _kept_workbook(name, data)
    if kept is None:                        # never a copy still holding the students: no copy at all
        _log.warning("%s: the Students or Combinations sheet could not be removed, the file was not kept", name)
        doubts = [*doubts, NOT_KEPT]
    body = {"merged": merged, "path": keep_file(db, name, kept) if kept is not None else None,
            "summary": summary, "doubts": doubts}
    if sizes:                               # kept beside the workbook: merged again with it at apply
        body["sizes_paths"] = [keep_file(db, n, d) for n, d in sizes]
    out = _plan_reading(db, name, body, "exact", tid)
    if out["status"] != "read":
        for p in _paths(body):
            _unlink(p)
    if alongside is not None and name.lower().endswith(".xlsx"):
        students = _sheet_students_reading(db, name, data, tid, merged)
        if students is not None:
            alongside.append(students)
        combinations = _sheet_combinations_reading(db, name, data, tid, merged)
        if combinations is not None:
            alongside.append(combinations)
    return out


def _sheet_students_reading(db, name: str, data: bytes, tid: str | None, plan: dict) -> dict | None:
    """The reading of a planning workbook's Students sheet (None when it holds no student), kept under
    its own file name: one of the same name replaces the plan reading's. The students are never in the
    plan reading's file or in anything the AI is shown."""
    try:
        if sheet_students(name, data) is None:
            return None
    except Exception as e:
        _log.warning("checking %s for a Students sheet failed (%s)", name, type(e).__name__)
        return None
    return _students_branch(db, f"{name} (Students sheet)", data, tid, plan, table=lambda n, d: sheet_students(name, d))


def _retry(db, name: str, data: bytes, message: str, tid: str | None) -> dict:
    return _failed(db, name, message, tid, path=keep_file(db, name, data), status="retry")


def _no_model(name: str) -> str:
    return f"The AI could not read {name} just now: check the model in Settings or try again."


def _smart(db, provider_factory, name: str, data: bytes, tid: str | None) -> dict:
    book = RR.csv_as_workbook(name, data) if name.lower().endswith(".csv") else data
    try:
        reading = writer.read(provider_factory, book, db.get_value("plan", tid), db=db)
    except R.ReadingError:                  # before RecipeError, its base: the run, not the recipes
        return _retry(db, name, data, RR.BUSY, tid)
    except writer.NoProvider as e:          # our own words: shown as they are
        return _retry(db, name, data, str(e), tid)
    except ProviderError as e:              # the provider's own words may be raw: logged, not shown
        _log.warning("reading %s: the provider failed: %s", name, e)
        return _retry(db, name, data, _no_model(name), tid)
    except (R.RecipeError, WorkbookError) as e:
        return _failed(db, name, f"{name} could not be read: {e}", tid)
    body = {**reading, "summary": reading.get("summary") or name}
    return _plan_reading(db, name, body, "saved" if reading.get("reader") else "ai", tid, data)


def _document(db, provider_factory, i: int, name: str, data: bytes, tid: str | None, timetables: list) -> dict | None:
    """A Word or PDF document: a plan reading stored now, or a timetable body put on `timetables` (stored
    once the drop's plan material is known; None returned then)."""
    try:
        ex = extract(name, data)
    except UnsupportedFile as e:
        return _failed(db, name, str(e), tid)
    if ex.kind == "pdf" and ex.warning:     # a scan: no text to read
        return _failed(db, name, ex.warning, tid)
    if document_students(ex):               # before the provider is made: a student list never reaches the AI
        return _students_failed(db, name, "a Word or PDF student list can't be read exactly; give it as Excel or CSV", tid)
    try:
        body = document.read_document(provider_factory(), name, ex.text, db.get_value("plan", tid), db.get_settings(tid))
    except ProviderError as e:
        _log.warning("reading %s: the provider failed: %s", name, e)
        return _retry(db, name, data, str(e) if isinstance(e, writer.NoProvider) else _no_model(name), tid)
    except document.DocumentError as e:
        return _failed(db, name, str(e), tid)
    if body.get("kind") == "timetable":
        timetables.append(([i], {"summary": body.get("summary") or name, "doubts": body.get("doubts") or [],
                                  "organisation": body["organisation"], "paths": [keep_file(db, name, data)]}))
        return None
    return _plan_reading(db, name, body, "ai", tid, data)


def asc_read(datas: list[bytes], mode: str = "keep", classes_as_planes: bool = False) -> tuple[dict, list[dict], dict | None]:
    """The draft an aSc export's PDFs make, validated like every other draft, and the week the export's grid
    proposes (None when it makes no week the app can keep). The week never stays on the draft."""
    pages: list = []
    for d in datas:
        pages.extend(asc_import.read_pdf_words(d))
    org, notes = asc_import.build_organisation(pages, mode=mode, classes_as_planes=classes_as_planes)
    week = org.pop("_week", None)
    return apply_org_patch(org, {}), notes, week


def asc_org(datas: list[bytes], mode: str = "keep", classes_as_planes: bool = False) -> tuple[dict, list[dict]]:
    """The draft an aSc export's PDFs make, validated like every other draft."""
    org, notes, _week = asc_read(datas, mode, classes_as_planes)
    return org, notes


def _same_week(a: dict, b: dict) -> bool:
    return (a["n_days"], a["spd"], a["open_per_day"]) == (b["n_days"], b["spd"], b["open_per_day"])


def _asc(db, asc: list[tuple[int, str, bytes]], tid: str | None, results: dict, timetables: list) -> None:
    try:
        org, notes, proposed = asc_read([d for _, _, d in asc], **OPTIONS)
    except asc_import.AscFormatError as e:
        for i, name, _ in asc:
            results[i] = _failed(db, name, str(e), tid)
        return
    except IntakeError as e:
        for i, name, _ in asc:
            results[i] = _failed(db, name, f"The timetable export could not be turned into a draft: {e}", tid)
        return
    doubts = [str(n.get("note")) for n in notes if isinstance(n, dict) and n.get("note")][:document.DOUBTS_MAX]
    body = {"summary": "Timetable export", "doubts": doubts, "_org": org,
            "paths": [keep_file(db, name, d) for _, name, d in asc]}
    if proposed is not None:                    # kept whole: whether it differs is asked afresh each time the card is shown
        body["week"] = proposed
    timetables.append(([i for i, _, _ in asc], body))


# ---------------------------------------------------------------------------
# student lists (spec docs/superpowers/specs/2026-10-16-students-design.md §2 Readers, §4): read
# exactly by students.read_rows, never by the AI. Every text the reading carries (summary, doubts,
# error, the upload's message) is made of counts, row numbers, IDs, class and option codes and the
# file's header words: never a name, never a raw cell (a problem's `detail` stays off the reading).
# ---------------------------------------------------------------------------

STUDENTS_SHEET = "students"
COMBINATIONS_SHEET = "combinations"
KEPT_OUT = (STUDENTS_SHEET, COMBINATIONS_SHEET)       # sheets read at the drop and left out of the kept copy and the AI's
NOT_AI = "Student lists are never sent to the AI reader."
RENAME = "Rename the columns to Student ID, Name, Class, Subject 1, Subject 2 and on, and give it again"
_ID_HEADS = ("student id", "admission no", "id", "reg no")
_NAME_HEADS = ("name", "student name", "full name")
_CLASS_HEADS = ("class", "form class", "form", "class code")
_SUBJECT_HEAD = re.compile(r"(?:subject|option|band) ?\d+")
_MERIT_HEADS = ("merit", "merit rank")      # never "rank" or "order": a B1 list may have a class rank
_CHOICE_HEAD = re.compile(r"(?:choice|preference|combination) ?(\d+)")         # ranked combination choices, never takes
_SUBJECT_LIST = "subjects"
_LIST_SPLIT = re.compile(r"[,;]")
_WANTED = {"id": "a student ID column (Student ID, Admission No, ID or Reg No)",
           "name": "a name column (Name, Student Name or Full Name)",
           "class": "a class column (Class, Form Class, Form or Class Code)",
           "subjects": "subject columns (Subject 1, Option 1, Band 1 and on, or Subjects) or Choice columns "
                       "(Choice 1, Preference 1, Combination 1 and on)"}
# the refusal net, wider than the reader: a row with one cell naming a class and another naming a
# student looks like a student list in other words, and is refused (never sent to the AI)
_LOOSE_CLASS = re.compile(r"\b(?:class(?:es)?|form|tutor ?group|home ?room|register)\b")
_LOOSE_WHO = re.compile(r"\b(?:name|surname|first ?name|id|admission|adm|reg|index|student|pupil)\b")
_NOT_WHO = re.compile(r"\b(?:teachers?|staff|subjects?|rooms?|venues?|department|dept|course|block|building)\b")
_PLAN_ROW = re.compile(r"\b(?:periods?|lessons?|load|hours)\b")
_PLAN_LINE = re.compile(r"\b(?:periods?|lessons?|load|hours|teachers?|staff|rooms?|venues?)\b")
LINE_WORDS = 12             # a PDF line longer than this is prose, not a header


def _student_columns(header: list) -> dict | None:
    """Where a header row's student-list columns are: {id, name, class, merit: index | None, subjects:
    [(index, label, a comma list?)], choices: [(index, label)] in choice-number order}; None unless it
    names a class and an ID or a name (it then looks like a student list, read or refused, and never
    goes to the AI)."""
    heads = [_head(c) for c in header]

    def first(names) -> int | None:
        return next((i for i, h in enumerate(heads) if h in names), None)
    cols = {"id": first(_ID_HEADS), "name": first(_NAME_HEADS), "class": first(_CLASS_HEADS),
            "subjects": [(i, " ".join(str(header[i]).split()), h == _SUBJECT_LIST) for i, h in enumerate(heads)
                         if h == _SUBJECT_LIST or _SUBJECT_HEAD.fullmatch(h)]}
    cols["merit"] = first(_MERIT_HEADS)
    ranked = sorted((int(m.group(1)), i) for i, h in enumerate(heads) if (m := _CHOICE_HEAD.fullmatch(h)))
    cols["choices"] = [(i, " ".join(str(header[i]).split())) for n, i in ranked if n >= 1]
    if cols["class"] is None or (cols["id"] is None and cols["name"] is None):
        return None
    return cols


def looks_like_students(cells: list, words: bool = False) -> bool:
    """A header row that looks like a student list's in other words: one cell naming a class (class, form,
    tutor group, homeroom, register) and another naming a student (name, surname, first name, ID,
    admission, adm, reg, index, student, pupil), as whole words in any case. A who-cell about a teacher,
    a subject or a room does not count, nor does a row with periods, lessons, load or hours (a plan's
    table). `words`: the cells are a text line's words (a PDF), which must be short and name no teacher."""
    heads = [_head(c) for c in cells]
    heads = [h for h in heads if h]
    if not heads or (words and len(heads) > LINE_WORDS):
        return False
    if any((_PLAN_LINE if words else _PLAN_ROW).search(h) for h in heads):
        return False
    klass = {i for i, h in enumerate(heads) if _LOOSE_CLASS.search(h)}
    who = {i for i, h in enumerate(heads) if _LOOSE_WHO.search(h) and not _NOT_WHO.search(h)}
    return any(i != j for i in klass for j in who)


def _find(rows: list[list], named: bool) -> tuple | None:
    """In a sheet's or a CSV's rows (searched to the end: a school's export may open with a title block):
    ("table", header, [(file row number, row)] under it) at the first row naming the student-list
    columns; else ("loose",) when some row looks like a student list's header in other words, or the
    sheet is `named` Students; else None."""
    loose = named
    for i, row in enumerate(rows):
        cells = list(row)
        if _student_columns(cells) is not None:
            return "table", cells, [(n, list(r)) for n, r in enumerate(rows[i + 1:], i + 2)]
        loose = loose or looks_like_students(cells)
    return ("loose",) if loose else None


def _csv_table(data: bytes) -> tuple | None:
    """A CSV read with a comma, then the delimiter the sniffer sees, then each other common one, until
    one finds a student list; a CSV none of them can split is looked at line by line for the loose test."""
    text = _decoded(data)
    for delimiter in _delimiters(text):
        try:
            found = _find(list(csv.reader(io.StringIO(text), delimiter=delimiter)), False)
        except csv.Error:
            continue
        if found is not None:
            return found
    if any(looks_like_students(re.split(r"[,;\t|]", line)) for line in text.splitlines()):
        return ("loose",)
    return None


def students_table(name: str, data: bytes) -> tuple | None:
    """What a file holds as a student list: ("table", header, numbered rows) to read exactly, ("loose",)
    to refuse, or None for any other file. An .xlsx's sheet named Students comes first, then its first
    sheet with the student-list columns, then any sheet that looks like a student list in other words.
    A workbook openpyxl cannot open is not one (whatever reads it next says why)."""
    lower = name.lower()
    if lower.endswith(".csv"):
        return _csv_table(data)
    if not lower.endswith(".xlsx"):
        return None
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        tables, loose = [], None
        for ws in wb.worksheets:
            named = str(ws.title).strip().casefold() == STUDENTS_SHEET
            ws.reset_dimensions()           # a stale <dimension ref="A1"/> would hide every other cell
            found = _find([list(r) for r in ws.iter_rows(values_only=True)], named)
            if found is not None and named:
                return found
            if found is not None and found[0] == "table":
                tables.append(found)
            elif found is not None:
                loose = loose or found
        return tables[0] if tables else loose
    finally:
        wb.close()


def sheet_students(name: str, data: bytes) -> tuple | None:
    """What a planning workbook's sheet named Students holds, to be read as `students_table` reads a
    file: ("table", header, numbered rows) to read exactly, ("loose",) to refuse; None when the workbook
    has no such sheet or it holds no student (no cell, no row under its header, or only the start
    wizard's example row). Only that sheet is looked at: any other sheet of a planning workbook is a
    department's, a venue list and so on."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        ws = next((w for w in wb.worksheets if str(w.title).strip().casefold() == STUDENTS_SHEET), None)
        if ws is None:
            return None
        ws.reset_dimensions()
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not any(_cell(v) for r in rows for v in r):
        return None
    found = _find(rows, True)
    if found is None or found[0] != "table":
        return found
    _, header, body = found
    body = [(n, r) for n, r in body if any(_cell(v) for v in r)]
    cols = _student_columns(header)
    if not body:
        return None
    if len(body) == 1 and cols is not None and cols["id"] is not None and cols["name"] is not None:
        row = body[0][1]
        if (_cell(row[cols["id"]]) == ST.EXAMPLE_ID if cols["id"] < len(row) else False) and \
                (_cell(row[cols["name"]]) == ST.EXAMPLE_NAME if cols["name"] < len(row) else False):
            return None
    return "table", header, body


def document_students(ex) -> bool:
    """A Word or PDF text (as extract gives it) holding a student list's header: a Word table row
    ("| a | b |") or a short PDF line that looks like one."""
    for line in ex.text.splitlines():
        line = line.strip()
        if ex.kind == "docx" and line.startswith("|"):
            if looks_like_students([c.strip() for c in line.strip("|").split("|")]):
                return True
        elif ex.kind == "pdf" and line and not line.startswith("## "):
            if looks_like_students(line.split(), words=True):
                return True
    return False


def _student_rows(cols: dict, rows: list[tuple[int, list]], header: list) -> list[dict]:
    """The rows students.read_rows takes ({row, id, name, class, subjects, columns}, `row` the file's own
    row number), blank rows left out."""
    def at(row, i):
        return _cell(row[i]) if i is not None and i < len(row) else ""
    out = []
    for n, row in rows:
        if not any(_cell(v) for v in row):
            continue
        subjects, columns = [], []
        for i, label, is_list in cols["subjects"]:
            cells = [x.strip() for x in _LIST_SPLIT.split(at(row, i))] if is_list else [at(row, i)]
            subjects += cells
            columns += [label] * len(cells)
        entry = {"row": n, "id": at(row, cols["id"]), "name": at(row, cols["name"]),
                 "class": at(row, cols["class"]), "subjects": subjects, "columns": columns}
        if cols["merit"] is not None:           # only when the file has the column: other files read as they did
            entry["merit"] = at(row, cols["merit"])
            entry["merit_column"] = " ".join(str(header[cols["merit"]]).split())
        if cols["choices"]:
            entry["choices"] = [at(row, i) for i, _ in cols["choices"]]
        out.append(entry)
    return out


def _students_failed(db, name: str, why: str, tid: str | None, message: str | None = None) -> dict:
    message = message or f"This looks like a student list but {why}. {NOT_AI}"
    entry = _entry(db, name, {"kind": "students", "via": "exact", "status": "failed", "sections": {},
                              "doubts": [], "summary": "Student list", "error": message}, tid)
    return {"name": name, "status": "failed", "id": entry["id"], "message": message}


def _to_check(problems: list[dict], students: list[dict], rows: list[dict]) -> int:
    """How many students (or rows that made none) a list's problems name."""
    made = {s["id"] for s in students}
    id_at = {r["row"]: r["id"] for r in rows}
    keys = set()
    for p in problems:
        if "student" in p:
            keys.add(p["student"])
        else:
            sid = id_at.get(p.get("row"), "")
            keys.add(sid if sid in made and p["kind"] != "duplicate_id" else ("row", p.get("row")))
    return len(keys)


def _students(db, name: str, found: tuple, tid: str | None, plan: dict | None = None) -> dict:
    if found[0] != "table":
        return _students_failed(db, name, f"its columns weren't recognised. {RENAME}", tid)
    _, header, body = found
    cols = _student_columns(header)
    # a list of choices needs no Subject columns: students choosing a combination have no subjects yet
    missing = [_WANTED[k] for k in ("id", "name", "class", "subjects")
               if cols[k] in (None, []) and not (k == "subjects" and cols["choices"])]
    if missing:
        return _students_failed(db, name, "its columns weren't recognised: it needs "
                                + " and ".join(missing), tid)
    plan = plan if plan is not None else _current(db, tid)
    rows = _student_rows(cols, body, header)
    if not rows:
        return _students_failed(db, name, "it has no rows under its header", tid)
    read, problems = ST.read_rows(rows, plan)
    if not read:
        return _students_failed(db, name, f"none of its {len(rows)} rows has a student ID", tid)
    codes = {c["code"] for c in plan["classes"]}
    if not any(s["class"] in codes for s in read):
        return _students_failed(db, name, "none of its classes is in the plan" if codes
                                else "the plan has no classes yet: load the plan first", tid)
    stored = [s for s in db.get_value("students", tid) or [] if isinstance(s, dict)]
    before = {s.get("id"): s for s in stored}
    ids = {s["id"] for s in read}
    added = sum(1 for s in read if s["id"] not in before)
    changed = sum(1 for s in read if s["id"] in before and before[s["id"]] != s)
    removed = sum(1 for sid in before if sid not in ids)
    texts = list(dict.fromkeys(p["text"] for p in problems))          # `text` only: `detail` is a raw cell
    replaces = ([f"This list replaces the {len(before):,} student{'' if len(before) == 1 else 's'} stored; "
                 f"{removed:,} of them {'isn' if removed == 1 else 'aren'}'t in this file"] if before else [])
    doubts = replaces + texts[:document.DOUBTS_MAX]
    if len(texts) > document.DOUBTS_MAX:
        doubts.append(f"and {len(texts) - document.DOUBTS_MAX} more to check")
    n, k = len(read), _to_check(problems, read, rows)
    summary = f"{n:,} student{'' if n == 1 else 's'}, {k:,} to check"
    secs = {"students": {"count": n, "added": added, "changed": changed, "removed": removed, "samples": [],
                         "ticked": True}}
    entry = _entry(db, name, {"kind": "students", "via": "exact", "status": "waiting", "sections": secs,
                              "doubts": doubts, "summary": summary, "students": read}, tid)
    return {"name": name, "status": "read", "id": entry["id"], "message": f"Student list: {summary}"}


def _students_branch(db, name: str, data: bytes, tid: str | None, plan: dict | None = None,
                     table=None) -> dict | None:
    """A student list read or refused (a stored reading's result); None for any other file. Whatever goes
    wrong here is a students card, logged by its kind only: an exception's words may hold a cell.
    `plan`, the plan to read against (the stored one by default); `table`, the function that finds the
    list in the file (`students_table` by default)."""
    try:
        found = (table or students_table)(name, data)
    except Exception as e:
        _log.warning("checking %s for a student list failed (%s)", name, type(e).__name__)
        return _students_failed(db, name, "", tid, message=f"{name} could not be checked for a student list, "
                                f"so it was not read; save it again as .xlsx or CSV and give it again. {NOT_AI}")
    if found is None:
        return None
    try:
        return _students(db, name, found, tid, plan)
    except M.PlanError:
        raise                               # the stored plan itself does not read: said as for any file
    except Exception as e:
        _log.warning("reading the student list %s failed (%s)", name, type(e).__name__)
        return _students_failed(db, name, "it could not be read exactly; check its columns and rows", tid)


# ---------------------------------------------------------------------------
# a workbook's Combinations sheet (option blocks B2): read exactly, never by the AI. The reading is a
# kind of its own ("combinations"), kept under "<file> (Combinations sheet)" beside the file's other
# readings; Apply replaces each level's combinations in `blocks_setup`. What it carries (summary, doubts,
# error) is row numbers, levels, combination codes and the validation's own words: no student data.
# ---------------------------------------------------------------------------

COMBOS_NOT_AI = "A Combinations sheet is never sent to the AI reader."
COMBOS_WANTED = "a Year column, a Code column, a Places column and Subject columns (Subject 1, Subject 2 and on)"
_COMBO_YEAR_HEADS = ("year", "level", "year group")
_COMBO_CODE_HEADS = ("code", "combination code", "combination")
_COMBO_NAME_HEADS = ("name", "combination name")
_COMBO_PLACES_HEADS = ("places", "seats")
_COMBO_SUBJECT_HEAD = re.compile(r"subject ?\d+")
SHOWN_CODE = 40             # a cell quoted in a doubt is cut to this


def _combination_columns(header: list) -> dict | None:
    """Where a header row's Combinations columns are: {year, code, name, places: index | None, subjects:
    [index]}; None unless it names a year, a code and a places column and at least one Subject column."""
    heads = [_head(c) for c in header]

    def first(names) -> int | None:
        return next((i for i, h in enumerate(heads) if h in names), None)
    cols = {"year": first(_COMBO_YEAR_HEADS), "code": first(_COMBO_CODE_HEADS), "name": first(_COMBO_NAME_HEADS),
            "places": first(_COMBO_PLACES_HEADS),
            "subjects": [i for i, h in enumerate(heads) if _COMBO_SUBJECT_HEAD.fullmatch(h)]}
    if None in (cols["year"], cols["code"], cols["places"]) or not cols["subjects"]:
        return None
    return cols


def _combinations_sheet(data: bytes) -> tuple | None:
    """A workbook's sheet named Combinations, looked at alone: ("table", header, numbered rows), ("loose",)
    when it has cells but no header naming the columns, ("example",) when its one row is the start wizard's
    example row, ("empty",) when it has no cell or no row under its header; None when there is no such sheet
    (or the workbook does not open)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        ws = next((w for w in wb.worksheets if " ".join(str(w.title).split()).casefold() == COMBINATIONS_SHEET), None)
        if ws is None:
            return None
        ws.reset_dimensions()               # a stale <dimension ref="A1"/> would hide every other cell
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not any(_cell(v) for r in rows for v in r):
        return ("empty",)
    at = next((i for i, r in enumerate(rows) if _combination_columns(r) is not None), None)
    if at is None:
        return ("loose",)
    header = rows[at]
    body = [(n, r) for n, r in enumerate(rows[at + 1:], start=at + 2) if any(_cell(v) for v in r)]
    if not body:
        return ("empty",)
    cols = _combination_columns(header)
    if len(body) == 1 and cols["name"] is not None:
        row = body[0][1]

        def cell(i):
            return _cell(row[i]) if i < len(row) else ""
        if cell(cols["code"]) == BK.EXAMPLE_CODE and cell(cols["name"]) == BK.EXAMPLE_NAME:
            return ("example",)
    return "table", header, body


def sheet_combinations(data: bytes) -> tuple | None:
    """What a workbook's sheet named Combinations holds: ("table", header, numbered rows) to read, ("loose",)
    when it has cells but no header naming the columns; None when the workbook has no such sheet or it
    holds nothing (no cell, no row under its header, or only the start wizard's example row). Only that
    sheet is looked at."""
    found = _combinations_sheet(data)
    return None if found is None or found[0] in ("empty", "example") else found


def holds_only_the_example(data: bytes) -> bool:
    """Whether the workbook's Combinations sheet holds the start wizard's example row and nothing else."""
    try:
        return _combinations_sheet(data) == ("example",)
    except Exception:
        return False


def _combinations_failed(db, name: str, why: str, tid: str | None, message: str | None = None) -> dict:
    message = message or f"This looks like a Combinations sheet but {why}. {COMBOS_NOT_AI}"
    entry = _entry(db, name, {"kind": "combinations", "via": "exact", "status": "failed", "sections": {},
                              "doubts": [], "summary": "Combinations", "error": message}, tid)
    return {"name": name, "status": "failed", "id": entry["id"], "message": message}


def _shown(text: str) -> str:
    return text if len(text) <= SHOWN_CODE else text[:SHOWN_CODE] + "..."


def _places(text: str):
    """A Places cell for the setup's own validation: a whole number as an int, nothing as None, anything else
    as the text (which it refuses)."""
    if not text:
        return None
    return int(text) if text.isdigit() else text


def _stored_students(db, tid: str | None) -> list:
    stored = db.get_value("students", tid)
    return stored if isinstance(stored, list) else []


def _remap_doubts(setup: dict, combos: list, students: list, plan: dict, key: str) -> list[str]:
    """A doubt for each stored code these combinations would make name other subjects while a student's
    choices or an override use it (R23): Apply refuses it. Codes and counts, never a student."""
    return [PM.remap_doubt(item) for item in PM.remapped_codes(
        setup.get("combinations") or [], combos, students, plan, key, setup.get("overrides"))]


def _stored_setup(stored: dict, key: str, doubts: list[str]) -> dict:
    """The setup stored for year key `key`, under its canonical key or an old raw one (R15); when two stored
    keys of the level differ (R20) the clash is added to `doubts` (Apply refuses it) and nothing is counted."""
    try:
        return BK.find_setup(stored, key)[1] or {}
    except BK.SetupClash as e:
        doubts.append(str(e))
        return {}


def _combination_rows(cols: dict, rows: list[tuple[int, list]], plan: dict) -> tuple[dict, list[str], int]:
    """({year key: [combination]} in sheet order, doubts, how many rows gave a doubt). A row is a doubt, with
    its row number, and is not kept, when its level is blank or not in the plan, or the setup's own
    validation (blocks.check_combinations) refuses it beside the rows kept before it."""
    levels = {ST.level_of(plan, c["code"]) for c in plan.get("classes") or []} - {""}
    known = ", ".join(sorted(levels, key=ST._natural)) or "none"

    def at(row, i):
        return _cell(row[i]) if i is not None and i < len(row) else ""
    kept: dict[str, list[dict]] = {}
    doubts, bad = [], 0
    for n, row in rows:
        level = at(row, cols["year"])
        subjects = [x for x in (at(row, i) for i in cols["subjects"])]
        candidate = {"code": at(row, cols["code"]), "name": at(row, cols["name"]), "subjects": subjects,
                     "places": _places(at(row, cols["places"]))}
        if not level:
            why = "it has no Year"
        elif level not in levels:
            why = f"level {_shown(level)} isn't in the plan (the plan's levels: {known})"
        else:
            key = f"level:{level}"
            try:
                BK.check_combinations([*kept.get(key, []), candidate])
            except ValueError as e:
                why = str(e)
            else:
                kept.setdefault(key, []).append(candidate)
                continue
        bad += 1
        doubts.append(f"Combinations row {n}: {why}")
    return ({key: BK.check_combinations(value) for key, value in kept.items()}, doubts, bad)


def _combinations(db, name: str, found: tuple, tid: str | None, plan: dict) -> dict:
    if found[0] != "table":
        return _combinations_failed(db, name, f"its columns weren't recognised: it needs {COMBOS_WANTED}", tid)
    _, header, body = found
    cols = _combination_columns(header)
    read, doubts, bad = _combination_rows(cols, body, plan)
    total = sum(len(v) for v in read.values())
    if not read:
        return _combinations_failed(db, name, "none of its rows could be read: " + "; ".join(doubts[:document.DOUBTS_MAX]), tid)
    stored = db.get_value("blocks_setup", tid)
    stored = stored if isinstance(stored, dict) else {}
    added = changed = removed = 0
    replaces = []
    students = _stored_students(db, tid)
    for key, combos in read.items():
        setup = _stored_setup(stored, key, replaces)
        replaces += _remap_doubts(setup, combos, students, plan, key)
        old = {c["code"]: c for c in (setup.get("combinations") or []) if isinstance(c, dict)}
        new = {c["code"]: c for c in combos}
        added += sum(1 for c in new if c not in old)
        changed += sum(1 for c in new if c in old and old[c] != new[c])
        removed += sum(1 for c in old if c not in new)
        if old:
            replaces.append(f"This replaces the {len(old):,} combination{'' if len(old) == 1 else 's'} stored for "
                            f"level {key.partition(':')[2]}")
    doubts = [*replaces, *doubts]
    if len(doubts) > document.DOUBTS_MAX:
        doubts = [*doubts[:document.DOUBTS_MAX], f"and {len(doubts) - document.DOUBTS_MAX} more to check"]
    m = len(read)
    summary = f"{total:,} combination{'' if total == 1 else 's'} at {m:,} level{'' if m == 1 else 's'}, {bad:,} to check"
    secs = {"combinations": {"count": total, "added": added, "changed": changed, "removed": removed,
                             "samples": [], "ticked": True}}
    entry = _entry(db, name, {"kind": "combinations", "via": "exact", "status": "waiting", "sections": secs,
                              "doubts": doubts, "summary": summary, "combinations": read}, tid)
    return {"name": name, "status": "read", "id": entry["id"], "message": f"Combinations: {summary}"}


def _sheet_combinations_reading(db, name: str, data: bytes, tid: str | None, plan: dict) -> dict | None:
    """The reading of a workbook's Combinations sheet (None when it holds nothing), kept under its own file
    name: one of the same name replaces the plan reading's. Whatever goes wrong is a combinations card,
    logged by its kind only: an exception's words may hold a cell."""
    label = f"{name} (Combinations sheet)"
    try:
        found = sheet_combinations(data)
    except Exception as e:
        _log.warning("checking %s for a Combinations sheet failed (%s)", name, type(e).__name__)
        return None
    if found is None:
        return None
    try:
        return _combinations(db, label, found, tid, plan)
    except Exception as e:
        _log.warning("reading the Combinations sheet of %s failed (%s)", name, type(e).__name__)
        return _combinations_failed(db, label, "it could not be read exactly; check its columns and rows", tid)


# ---------------------------------------------------------------------------
# the five CSV files (csv-pack spec docs/superpowers/specs/2026-10-20-csv-pack-design.md §2 and its
# "Calls made by the designer" §2-§4): staff, venues, requirements, combinations and the level students
# list, each recognised by its headers and read exactly by csvpack, never by the AI. A drop's pack files
# are read in dependency order (PACK_ORDER), each against what is stored overlaid with the pack files
# read before it in the same drop, so the upload order does not matter. A card carries counts and the
# readers' problems (file, row and header words; never a student's name), and "N rows not applied" for
# the rows a structural problem left out. Logs carry type names only.
# ---------------------------------------------------------------------------

PACK_ORDER = ("requirements", "combinations", "staff", "venues", "level_students")
PACK_NOT_AI = "These files are read exactly and never sent to the AI reader."
_PACK_TITLE = {"staff": "Staff", "venues": "Rooms", "requirements": "Requirements", "combinations": "Combinations",
               "level_students": "Students"}
_PACK_READ = {"staff": CP.read_staff, "venues": CP.read_venues, "requirements": CP.read_requirements,
              "combinations": CP.read_combinations, "level_students": CP.read_level_students}
NOT_IN_FILE_SHOWN = 8       # names listed in a "not in this file" doubt


def _new_drop() -> dict:
    """What the pack files of one drop have read so far, for the files read after them."""
    return {"levels": {}, "combinations": {}, "staff": []}


def _pack_found(name: str, data: bytes) -> tuple | None:
    """(kind, header, numbered rows) of a pack file; None for any other file. A workbook with a Students
    or Combinations sheet is B1's or B2's, and a header with a Class column (B1's list, a home-room list, a
    per-class table) is not one of the five: it keeps the route it had.
    A classless header naming students that `detect` does not take is still the level students list
    (refused with the column to add, never given to the AI)."""
    lower = name.lower()
    if not lower.endswith((".csv", ".xlsx")):
        return None
    if lower.endswith(".xlsx") and _names_students_sheet(data):
        return None
    found = CP.table(name, data)
    if found is None:
        return None
    header, rows = found
    if _student_columns(header) is not None or CP.names_a_class(header):
        return None
    kind = CP.detect(header) or ("level_students" if CP.names_students(header) else None)
    return None if kind is None else (kind, header, rows)


def _pack_kind(name: str, data: bytes) -> str | None:
    """The pack kind of a file, for the order a drop is read in; None when it is none (or cannot be told)."""
    try:
        found = _pack_found(name, data)
    except Exception:                       # its own card says what is wrong with it
        return None
    return found[0] if found else None


def _stored_combinations(db, tid: str | None) -> dict:
    """{level_key: [combination]} of each level's option-block setup stored. A level's own setup, the one
    combinations.csv replaces ("level:3" or an old spelling of it, `blocks.find_setup`), comes before a
    stream's ("level:3E", R26), which stands in only when the level has none."""
    stored = db.get_value("blocks_setup", tid)
    stored = stored if isinstance(stored, dict) else {}

    def combos_of(setup) -> list:
        return [c for c in (setup or {}).get("combinations") or [] if isinstance(c, dict) and c.get("code")]
    out: dict = {}
    for key in sorted(k for k in stored if isinstance(k, str)):
        if key.startswith("level:") and isinstance(stored[key], dict) and combos_of(stored[key]):
            out.setdefault(SP.level_key(key.partition(":")[2]), combos_of(stored[key]))
    for level in out:
        try:
            own = combos_of(BK.find_setup(stored, f"level:{level}")[1])
        except BK.SetupClash:               # the card says so (_stored_setup); Apply refuses it
            continue
        if own:
            out[level] = own
    return out


def _pack_context(db, tid: str | None, drop: dict) -> dict:
    """The facts a pack file is checked against: what is stored, overlaid with the drop's files read so far."""
    context = {"levels": {**LV.get(db, tid), **drop["levels"]},
               "combinations": {**_stored_combinations(db, tid), **drop["combinations"]},
               "plan": _current(db, tid)}
    if drop["staff"]:
        context["staff"] = list(drop["staff"])
    return context


def _pack_failed(db, name: str, kind: str, message: str, tid: str | None) -> dict:
    entry = _entry(db, name, {"kind": kind, "via": "exact", "status": "failed", "sections": {}, "doubts": [],
                              "summary": _PACK_TITLE[kind], "error": message}, tid)
    return {"name": name, "status": "failed", "id": entry["id"], "message": message}


def _plural(n: int, word: str, many: str | None = None) -> str:
    return f"{n:,} {word if n == 1 else many or word + 's'}"


def _not_in_file(names: list, what: str) -> list[str]:
    """The "not in this file" doubt of a staff or venues card: they are kept, never deleted (Calls §2)."""
    if not names:
        return []
    shown = ", ".join(str(n) for n in names[:NOT_IN_FILE_SHOWN])
    more = f" and {len(names) - NOT_IN_FILE_SHOWN} more" if len(names) > NOT_IN_FILE_SHOWN else ""
    return [f"{_plural(len(names), what)} in the plan {'is' if len(names) == 1 else 'are'} not in this file "
            f"and {'is' if len(names) == 1 else 'are'} kept: {shown}{more}"]


def _pack_card(kind: str, items, context: dict, db, tid: str | None) -> tuple[dict, list[str], str, dict]:
    """(sections, the doubts that come before the problems, the summary's count words, what the reading
    keeps) of a pack file's items."""
    plan = context["plan"]
    if kind == "staff":
        _, n = PM.merge_staff(plan["staff"], items)
        sec = {"count": len(items), "added": n["added"], "changed": n["changed"]}
        return {"staff": sec}, _not_in_file(n["missing"], "teacher"), _plural(len(items), "teacher"), \
            {"items": items}
    if kind == "venues":
        _, n = PM.merge_venues(plan["venues"], items)
        sec = {"count": len(items), "added": n["added"], "changed": n["changed"]}
        return {"venues": sec}, _not_in_file(n["missing"], "room"), _plural(len(items), "room"), {"items": items}
    if kind == "requirements":
        _, n = PM.merge_levels(LV.get(db, tid), items)
        rows = sum(len(lv["subjects"]) for lv in items.values())
        sec = {"count": rows, "added": n["added"], "changed": n["changed"]}
        return ({"subjects": sec}, [], f"{_plural(rows, 'subject')} at {_plural(len(items), 'level')}",
                {"levels": items})
    if kind == "combinations":
        stored = db.get_value("blocks_setup", tid)
        stored = stored if isinstance(stored, dict) else {}
        found = {f"level:{key}": combos for key, combos in items.items()}
        added = changed = removed = 0
        replaces = []
        students = _stored_students(db, tid)
        for key, combos in found.items():
            setup = _stored_setup(stored, key, replaces)
            old = {c["code"]: c for c in (setup.get("combinations") or []) if isinstance(c, dict)}
            new = {c["code"]: c for c in combos}
            added += sum(1 for c in new if c not in old)
            changed += sum(1 for c in new if c in old and old[c] != new[c])
            removed += sum(1 for c in old if c not in new)
            if old:
                replaces.append(f"This replaces the {_plural(len(old), 'combination')} stored for "
                                f"level {key.partition(':')[2]}")
            replaces += _remap_doubts(setup, combos, students, plan, key)
        total = sum(len(v) for v in found.values())
        sec = {"count": total, "added": added, "changed": changed, "removed": removed}
        return ({"combinations": sec}, replaces, f"{_plural(total, 'combination')} at {_plural(len(found), 'level')}",
                {"combinations": found})
    stored = _stored_students(db, tid)
    _, n = PM.merge_students(stored, items, plan)
    sec = {"count": len(items), "added": n["added"], "changed": n["changed"], "removed": n["removed"]}
    replaces = ([f"This list replaces the students stored at its levels; {_plural(n['removed'], 'student')} "
                 f"there {'isn' if n['removed'] == 1 else 'aren'}'t in this file"] if n["removed"] else [])
    ids = {s["id"] for s in items}
    classed = sum(1 for s in stored if isinstance(s, dict) and s.get("id") in ids and ST.has_class(s))
    if classed:                             # R27(d): a count, never a name
        replaces.append(f"{classed:,} of these students {'has' if classed == 1 else 'have'} a class now; "
                        f"Apply removes it")
    return {"students": sec}, replaces, _plural(len(items), "student"), {"students": items}


def _pack(db, name: str, found: tuple, tid: str | None, drop: dict) -> dict:
    kind, header, rows = found
    lacking = CP.missing(kind, header)
    if lacking:
        return _pack_failed(db, name, kind, f"This looks like {CP.FILES[kind]} but it needs a "
                            f"{' and a '.join(lacking)} column: add {'it' if len(lacking) == 1 else 'them'} "
                            f"and give it again. {PACK_NOT_AI}", tid)
    context = _pack_context(db, tid, drop)
    items, problems = _PACK_READ[kind](rows, header, context, file=name)
    if kind == "level_students":            # B1's shape: a student with a level and no class takes nothing yet
        items = [{**s, "takes": [], "placed_by_app": []} for s in items]
    if not items and not problems and not rows and getattr(rows, "examples", 0):
        return _pack_failed(db, name, kind, f"{name} holds only the example row: add your rows and give it "
                                            f"again. {PACK_NOT_AI}", tid)
    if not items:
        why = "; ".join(p["text"] for p in problems[:3]) or f"{name} has no rows under its header"
        return _pack_failed(db, name, kind, f"Nothing in {name} could be read: {why}. {PACK_NOT_AI}", tid)
    secs, before, counted, kept = _pack_card(kind, items, context, db, tid)
    for s in secs.values():
        s.update({"samples": [], "ticked": True})
    dropped = {p["row"] for p in problems if p["dropped"] and p["row"] is not None}
    check = {p["row"] for p in problems if not p["dropped"] and p["row"] is not None} - dropped
    texts = list(dict.fromkeys(p["text"] for p in problems))
    doubts = [*before, *texts]
    if len(doubts) > document.DOUBTS_MAX:
        doubts = [*doubts[:document.DOUBTS_MAX], f"and {len(doubts) - document.DOUBTS_MAX} more to check"]
    examples = getattr(rows, "examples", 0)
    summary = ", ".join([counted, *([f"{_plural(len(dropped), 'row')} not applied"] if dropped else []),
                         *([f"{len(check):,} to check"] if check else []),
                         *([f"{_plural(examples, 'example row')} skipped"] if examples else [])])
    entry = _entry(db, name, {"kind": kind, "via": "exact", "status": "waiting", "sections": secs,
                              "doubts": doubts, "summary": summary, **kept}, tid)
    if kind == "requirements":              # the files read after it in the drop see these levels
        drop["levels"].update(items)
    elif kind == "combinations":
        drop["combinations"].update(items)
    elif kind == "staff":                   # read ahead already (_preread_staff) unless given on its own
        known = {ST.norm(s["name"]) for s in drop["staff"]}
        drop["staff"] += [s for s in items if ST.norm(s["name"]) not in known]
    return {"name": name, "status": "read", "id": entry["id"], "message": f"{_PACK_TITLE[kind]}: {summary}"}


SNIFF_LINES = 50            # rows looked at for a header naming students when a file cannot be checked


def _sniffed_rows(name: str, data: bytes):
    """The first rows of a .csv (split on any common delimiter) or of each sheet of an .xlsx, as cells."""
    if name.lower().endswith(".csv"):
        for line in _decoded(data).splitlines()[:SNIFF_LINES]:
            yield re.split(r"[,;\t|]", line)
        return
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=SNIFF_LINES, values_only=True):
                yield ["" if v is None else str(v) for v in row]
    finally:
        wb.close()


def _names_students(name: str, data: bytes) -> bool:
    """Whether a file that could not be checked for a pack file has a row naming students (a Student ID
    column, or a class beside a name or ID): such a file fails closed (R27(c)). True when even that cannot
    be told."""
    try:
        return any(CP.names_students(cells) or looks_like_students(cells) for cells in _sniffed_rows(name, data))
    except Exception:
        return True


def _pack_branch(db, name: str, data: bytes, tid: str | None, drop: dict | None = None) -> dict | None:
    """A pack file read or refused (a stored reading's result); None for any other file. Whatever goes
    wrong while reading it is its own failed card, logged by type name only: never the AI. A file that
    cannot even be checked is a failed card too when it names students (R27(c)); any other goes on."""
    try:
        found = _pack_found(name, data)
    except Exception as e:
        _log.warning("checking %s for a pack file failed (%s)", name, type(e).__name__)
        if _names_students(name, data):
            return _pack_failed(db, name, "level_students", f"{name} names students but could not be checked "
                                f"exactly, so it was not read; check its columns and give it again. {PACK_NOT_AI}",
                                tid)
        return None
    if found is None:
        return None
    try:
        return _pack(db, name, found, tid, drop if drop is not None else _new_drop())
    except M.PlanError:
        raise                               # the stored plan itself does not read: said as for any file
    except Exception as e:
        _log.warning("reading the %s file %s failed (%s)", found[0], name, type(e).__name__)
        return _pack_failed(db, name, found[0], f"{name} could not be read exactly; check its columns and rows. "
                                                f"{PACK_NOT_AI}", tid)


def _preread_staff(items: list[tuple[str, bytes]], positions) -> list[dict]:
    """The teachers of a drop's staff files, read before anything else: requirements.csv, read first,
    takes a subject's department from the one department whose staff list it (Calls §6). A teacher's
    fields do not depend on the other files; only their problems do, and those come when the file is read."""
    out = []
    for i in positions:
        name, data = items[i][0] or "upload", items[i][1]
        try:
            found = _pack_found(name, data)
            if found and found[0] == "staff" and not CP.missing("staff", found[1]):
                out += CP.read_staff(found[2], found[1], {}, file=name)[0]
        except Exception:                   # its own card says what is wrong with it
            continue
    return out


# ---------------------------------------------------------------------------
# a drop of files
# ---------------------------------------------------------------------------

def _one(db, provider_factory, i: int, name: str, data: bytes, tid: str | None, timetables: list,
         sizes: list[tuple[str, bytes]] = (), alongside: list | None = None, drop: dict | None = None) -> dict | None:
    lower = name.lower()
    if is_planning(name, data):
        return _exact(db, name, data, tid, sizes=sizes, alongside=alongside)
    if sizes_text(name, data) is not None:  # class sizes: read exactly, as the plan upload once took them
        return _exact(db, name, data, tid, "Class sizes")
    if lower.endswith((".xlsx", ".csv")):   # one of the five CSV files: read exactly or refused, never by the AI
        pack = _pack_branch(db, name, data, tid, drop)
        if pack is not None:
            return pack
    if lower.endswith((".xlsx", ".csv")):   # a student list: read exactly or refused, never by the AI
        students = _students_branch(db, name, data, tid)
        if students is not None:
            if lower.endswith(".xlsx"):     # a list with a Combinations sheet: that is read too, exactly
                combinations = _sheet_combinations_reading(db, name, data, tid, _current(db, tid))
                if combinations is not None and alongside is not None:
                    alongside.append(combinations)
                # (without `alongside` the combinations reading is still stored, and shows with the others)
            return students
    if lower.endswith(".xlsx"):             # a Combinations sheet is read exactly, and is not what the AI is given
        combinations = _sheet_combinations_reading(db, name, data, tid, _current(db, tid))
        if not _has_other_sheets(data):     # nothing else in the file for the AI: the exact reading is the whole result
            if combinations is not None:
                return combinations
            if holds_only_the_example(data):
                return _failed(db, name, f"{name} has nothing to read: its Combinations sheet holds only the "
                                         "example row", tid)
            return _failed(db, name, f"{name} has nothing to read.", tid)
        else:
            if combinations is not None and alongside is not None:
                alongside.append(combinations)
            data = without_students_sheet(data)
            if data is None:
                return _failed(db, name, f"{name} could not be read: its Combinations sheet couldn't be set aside "
                                         "from the AI reader", tid)
    if lower.endswith((".xlsx", ".csv")):
        return _smart(db, provider_factory, name, data, tid)
    if lower.endswith((".docx", ".pdf")):
        return _document(db, provider_factory, i, name, data, tid, timetables)
    if lower.endswith((".doc", ".xls")):
        try:
            extract(name, data)
        except UnsupportedFile as e:
            return _failed(db, name, str(e), tid)
    return _failed(db, name, f"{name} is not a file the Assistant reads: give it PDF, Word (.docx), Excel (.xlsx) or CSV.", tid)


def _folding(items: list[tuple[str, bytes]]) -> tuple[int | None, list[int]]:
    """(the position of the drop's first planning workbook, the positions of its sizes CSVs), or (None, [])
    when the drop holds no planning workbook: the sizes then go on their own cards."""
    def ok(i: int, name: str, data: bytes, test) -> bool:
        if i >= FILES_MAX or len(data) > MAX_UPLOAD:
            return False
        try:
            return bool(test(name or "upload", data))
        except Exception:                   # not readable here: the file's own card says why
            return False
    book = next((i for i, (n, d) in enumerate(items) if ok(i, n, d, is_planning)), None)
    if book is None:
        return None, []
    return book, [i for i, (n, d) in enumerate(items) if ok(i, n, d, lambda n, d: sizes_text(n, d) is not None)]


def read_files(db, provider_factory, items: list[tuple[str, bytes]], tid: str | None = None) -> list[dict]:
    """Blocking. Each file of `items` ((name, bytes), in order) as a stored reading of timetable `tid`
    (the working one when None); one result per file, {name, status: read | failed | retry, id, message}.
    Never raises: whatever goes wrong is that file's failed card, so one file never stops the others.
    A file past FILES_MAX gets a card only: no stored reading (it would push the drop's own readings out).
    A drop's sizes CSVs go into the reading of its first planning workbook (their sizes applied on top of
    its merge, as the plan upload once took them); each says so, with that reading's id."""
    tid = tid or db.working_timetable()     # pinned: the selection may change while the model reads
    try:
        sweep(db)
    except Exception:
        _log.exception("sweeping kept files failed")
    book, folded = _folding(items)
    # the five CSV files first, in dependency order (each read against those before it), then the rest
    # in upload order; the results still line up with the files
    rank = {}
    for i, (name, data) in enumerate(items[:FILES_MAX]):
        if len(data) <= MAX_UPLOAD and i not in folded:
            kind = _pack_kind(name or "upload", data)
            if kind is not None:
                rank[i] = PACK_ORDER.index(kind)
    sequence = sorted(range(len(items)), key=lambda i: (0, rank[i], i) if i in rank else (1, 0, i))
    drop = _new_drop()
    drop["staff"] = _preread_staff(items, [i for i in rank if rank[i] == PACK_ORDER.index("staff")])
    results: dict[int, dict] = {}
    also: list[dict] = []                   # a planning workbook's Students sheet: after every file's own result
    asc: list[tuple[int, str, bytes]] = []
    timetables: list = []                   # ([positions], body) of timetable readings, stored last
    for i in sequence:
        name, data = items[i]
        name = name or "upload"
        try:
            if i >= FILES_MAX:
                results[i] = {"name": name, "status": "failed",
                              "message": f"Up to {FILES_MAX} files at a time: give {name} again on its own."}
            elif len(data) > MAX_UPLOAD:
                results[i] = _failed(db, name, f"{name} is over 20 MB", tid)
            elif i in folded:
                continue                    # read with the workbook, after it
            elif name.lower().endswith(".pdf") and asc_import.is_asc_pdf(data):
                asc.append((i, name, data))
            else:
                sizes = [(items[j][0] or "upload", items[j][1]) for j in folded] if i == book else []
                out = _one(db, provider_factory, i, name, data, tid, timetables, sizes, also, drop)
                if out is not None:
                    results[i] = out
        except M.PlanError as e:            # the stored plan itself does not read
            results[i] = _failed(db, name, f"{RR.PLAN_UNREADABLE} ({e})", tid)
        except Exception:
            _log.exception("reading %s failed", name)
            results[i] = _failed(db, name, f"Could not read {name}; try again.", tid)
    for j in folded:
        name, data = items[j][0] or "upload", items[j][1]
        wb = results.get(book) or {}
        if wb.get("status") == "read":
            results[j] = {"name": name, "status": "read", "id": wb["id"],
                          "message": f"The sizes in {name} are in the {wb['name']} card."}
            continue
        try:                                # the workbook did not read: the sizes on a card of their own
            results[j] = _one(db, provider_factory, j, name, data, tid, timetables, alongside=also)
        except M.PlanError as e:
            results[j] = _failed(db, name, f"{RR.PLAN_UNREADABLE} ({e})", tid)
        except Exception:
            _log.exception("reading %s failed", name)
            results[j] = _failed(db, name, f"Could not read {name}; try again.", tid)
    if asc:
        try:
            _asc(db, asc, tid, results, timetables)
        except Exception:
            _log.exception("reading the timetable export failed")
            for i, name, _ in asc:
                results[i] = _failed(db, name, "Could not read the timetable export; try again.", tid)
    plan_material = any(r["status"] == "read" for r in results.values())
    for positions, body in timetables:
        names = [items[i][0] or "upload" for i in positions]
        try:
            out = _timetable_reading(db, ", ".join(names), body, not plan_material, tid)
        except Exception:
            _log.exception("storing the timetable reading failed")
            out = None
        for i, name in zip(positions, names):
            results[i] = ({**out, "name": name} if out is not None
                          else _failed(db, name, f"Could not read {name}; try again.", tid))
    return [*(results[i] for i in sorted(results)), *also]     # last: the results line up with the files


def retry(db, provider_factory, rid: str, tid: str | None = None) -> dict:
    """Blocking. The kept file of reading `rid`, waiting to be tried again, read once more."""
    tid = tid or db.working_timetable()
    entry = next((e for e in RR._stored(db, tid) if e.get("id") == rid), None)
    reading = (entry or {}).get("reading") or {}
    path = reading.get("path")
    if entry is None or reading.get("status") != "retry" or not path:
        return {"status": "failed", "message": "Nothing to try again."}
    try:
        data = Path(path).read_bytes()
    except OSError:
        return {"name": entry.get("file") or "", "status": "failed", "message": "The file is gone: give it again."}
    out = read_files(db, provider_factory, [(entry.get("file") or "upload", data)], tid)[0]
    if path not in _named(db):
        _unlink(path)
    return out


# ---------------------------------------------------------------------------
# what the page sees, and clearing
# ---------------------------------------------------------------------------

_OLD_SECTION = {"venues": "rooms", "homes": "classes"}


def _old_sections(r: dict) -> dict:
    samples = r.get("samples") or {}
    out: dict[str, dict] = {}
    for target, n in (r.get("counts") or {}).items():
        name = _OLD_SECTION.get(target, target)
        s = out.setdefault(name, {"count": 0, "added": 0, "changed": 0, "samples": [], "ticked": True})
        s["count"] += n
        s["added"] += n
        s["samples"] += [str(x.get("name") or x.get("code") or x.get("subject") or "")
                         for x in samples.get(target, []) if isinstance(x, dict)]
    return out


def view(entry: dict, settings: dict | None = None) -> dict:
    r = entry.get("reading") or {}
    old = "kind" not in r                   # stored before readings had kinds: a smart-reader plan reading
    return {"id": entry.get("id"), "file": entry.get("file") or "", "made_at": entry.get("made_at"),
            "kind": r.get("kind") or "plan",
            "via": ("saved" if r.get("reader") else "ai") if old else r.get("via") or "ai",
            "status": r.get("status") or "waiting",
            "summary": r.get("summary") or "",
            "sections": _old_sections(r) if old else r.get("sections") or {},
            "doubts": r.get("doubts") or [], "error": r.get("error") or "",
            "options": r.get("options"), "undo": bool(r.get("undo")),
            "applied": r.get("applied"),      # {section: count} once applied: what the card says was applied
            "week_adopted": bool((r.get("undo") or {}).get("week_adopted")),
            **_week_view(r, settings)}


def _week_view(r: dict, settings: dict | None) -> dict:
    """{"week": {"file": [lines], "mine": [lines], "time": the file's week}} when an export's week differs from
    the timetable's as it is now, as the card shows it; nothing otherwise (or without `settings`)."""
    if not (settings and isinstance(r.get("week"), dict)):
        return {}
    try:
        theirs = W.shape(W.derive(r["week"]))
    except (KeyError, TypeError, ValueError):
        return {}
    mine = W.shape(settings.get("time"), settings.get("rules"))
    if _same_week(mine, theirs):
        return {}
    return {"week": {"file": W.describe(theirs), "mine": W.describe(mine), "time": r["week"]}}


def clear(db, tid: str | None = None) -> None:
    """Drop the readings still waiting (or failed, or to try again) and their kept files."""
    tid = tid or db.working_timetable()
    gone: list[str] = []
    with RR._lock:
        kept = []
        for e in RR._stored(db, tid):
            if view(e)["status"] in WAITING:
                gone += _paths(e["reading"])
            else:
                kept.append(e)
        db.set_value(RR.KEY, kept, tid)
    for p in gone:
        _unlink(p)
