"""A drop of generic CSV files (spec docs/superpowers/specs/2026-10-25-per-template-csv-files-design.md, its
"Calls made"): staff.csv, venues.csv, units.csv and duties.csv, one per sheet of the generic workbook
(`importer.read_generic_workbook`), told apart from the five school files by their headers, not their names.
The drop is assembled into one in-memory workbook with a `CSV files` marker sheet and read exactly, as the
workbook would be; a sheet it lacks is taken from the stored plan (`read_generic_workbook`, at the drop and
again at Apply). Never the AI: a drop that cannot be read is refused with what to do.

`detect` names a header's kind; `drop_set` finds a drop's generic files; `assemble` builds the workbook;
`sort_out` decides what a drop comes to. Pure: no database, no logging of cell values."""
from __future__ import annotations

import csv
import io
import re

from ...config import MAX_UPLOAD
from . import csvpack as CP
from .tabular import decoded, delimiters

KINDS = ("staff", "venues", "units", "duties")              # the workbook's sheet order
REQUIRED = ("staff", "units", "duties")                     # venues.csv is optional, as the Venues sheet is
MARKER_SHEET = "CSV files"
SUMMARY = "CSV files"
TITLES = {"staff": "Staff", "venues": "Venues", "units": "Units", "duties": "Duties"}
# each sheet's columns as the workbook writes them (`wizard.instantiate.sheets_for`): a header that names one
# by `CP._key` ("Per-cycle", "LOAD FACTOR") is written as it, the way the importer looks it up
HEADS = {
    "staff": ("Name", "Role", "Availability", "Load factor"),
    "venues": ("Name", "Kind", "Capacity"),
    "units": ("Name", "Size"),
    "duties": ("Name", "Unit", "Per cycle", "Length", "Venue kind", "Who can do it", "Together with", "Co-staffed"),
}
_CANON = {kind: {CP._key(h): h for h in heads} for kind, heads in HEADS.items()}
_SECTION = {"staff": ("staff", "staff"), "units": ("classes", "units"), "duties": ("requirements", "duties")}
_WHOLE = re.compile(r"-?\d+")
_DECIMAL = re.compile(r"-?\d+\.\d+")
_DECIMAL_COMMA = re.compile(r"-?\d+,\d+")          # "0,5": a `;` file from a decimal-comma Excel (`CP._number`)
# the columns the importer reads as numbers: only there is "0,5" a decimal (elsewhere a comma is a list's)
_NUMERIC = {CP._key(h) for h in ("Load factor", "Capacity", "Size", "Per cycle", "Length", "Co-staffed")}


def detect(header) -> str | None:
    """The generic kind a header names, or None: Name with Unit and Per cycle or Length is duties; Name with
    Role and Availability is staff (spec "Calls made": a Role alone, as an address list has, is not); Name with Size and no Kind is units; Name with Kind and no Room is venues.
    A header one of the five school files has (`CP.detect`) is never generic."""
    header = list(header or [])
    if CP.detect(header) is not None:
        return None
    keys = {CP._key(h) for h in header}
    if "name" not in keys:
        return None
    if "unit" in keys and keys & {"per cycle", "length"}:
        return "duties"
    if {"role", "availability"} <= keys:
        return "staff"
    if "size" in keys and "kind" not in keys:
        return "units"
    if "kind" in keys and "room" not in keys:
        return "venues"
    return None


def _table(name: str, data: bytes):
    """(kind, header, rows) of a generic .csv (`CP.table`'s decoding and shaping, each delimiter tried until
    one gives a generic header), or None."""
    if not str(name or "").lower().endswith(".csv"):
        return None
    text = decoded(data)
    for delimiter in delimiters(text):
        try:
            shaped = CP._shape(enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), 1))
        except csv.Error:
            continue
        kind = detect(shaped[0]) if shaped is not None else None
        if kind is not None:
            return kind, shaped[0], shaped[1]
    return None


def kind_of(name: str, data: bytes) -> str | None:
    """A file's generic kind; None when it is none, or cannot be told."""
    try:
        found = _table(name, data)
    except Exception:                       # its own card says what is wrong with it
        return None
    return found[0] if found else None


def generic_files(items: list[tuple[str, bytes]]) -> dict[int, str]:
    """{position: kind} of every generic file of `items` within the upload limit, in drop order."""
    out = {}
    for i, (name, data) in enumerate(items):
        if len(data) <= MAX_UPLOAD:
            kind = kind_of(name or "upload", data)
            if kind is not None:
                out[i] = kind
    return out


def drop_set(items: list[tuple[str, bytes]]) -> dict[str, int]:
    """{kind: position} of a drop's generic files, the first of each kind; empty unless the drop holds a
    generic staff, units or duties file. A lone venues.csv is never a generic drop: it keeps its old route; nor,
    on a school's timetable, is a drop lacking one of those three (`inbox._generic_shaped`, the caller's)."""
    out: dict[str, int] = {}
    for i, kind in generic_files(items).items():
        out.setdefault(kind, i)
    return out if set(out) & set(REQUIRED) else {}


def _value(text: str, numeric: bool = False):
    """A cell as the workbook would hold it: a number typed as one (so the importer and its row hashes see what
    a filled-in workbook gives them), "" as an empty cell, anything else as written. A whole number keeps its
    text when it would not read back the same ("007"). In a `numeric` column one comma and no dot is a decimal
    comma ("0,5" is 0.5), as `CP._number` reads the school files."""
    if not text:
        return None
    if _WHOLE.fullmatch(text) and str(int(text)) == text:
        return int(text)
    if _DECIMAL.fullmatch(text):
        return float(text)
    if numeric and _DECIMAL_COMMA.fullmatch(text):
        return float(text.replace(",", "."))
    return text


def _real_rows(name: str, data: bytes) -> int:
    found = _table(name, data)
    return len(found[2]) if found else 0


def assemble(files: dict[str, tuple[str, bytes]]) -> tuple[bytes, int]:
    """{kind: (name, bytes)} as one workbook: (its xlsx bytes, the number of real rows). One sheet per kind
    given, titled as the workbook's, holding the header and the rows that are not examples; then the
    `CSV files` sheet listing the kinds given, one a row, which tells the importer a sheet it lacks is the
    stored plan's."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    count = 0
    given = [kind for kind in KINDS if kind in files]
    for kind in given:
        found = _table(*files[kind])
        if found is None:                   # no sheet at all: an empty one would wipe the stored section
            continue
        _, header, rows = found
        ws = wb.create_sheet(TITLES[kind])
        ws.append([_CANON[kind].get(CP._key(h), h) for h in header])
        numeric = [CP._key(h) in _NUMERIC for h in header]
        for _, cells in rows:               # cells past the header are refused first (`sort_out`), never kept
            ws.append([_value(c, n) for c, n in zip(cells[:len(header)], numeric)])
        count += len(rows)
    marker = wb.create_sheet(MARKER_SHEET)
    for kind in given:
        marker.append([kind])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), count


def _overflowing(name: str, data: bytes) -> list[int]:
    """The file's own numbers of the rows with cells past the header (an unquoted list): the importer reads
    by header, so those cells would be lost, never read."""
    found = _table(name, data)
    return [n for n, cells in found[2] if len(cells) > len(found[1])] if found else []


OVERFLOW = "more cells than the header — put a list in quotes"   # the school files' wording (`CP._Out.overflows`)
SHOWN_ROWS = 3              # rows named on a refused card


def _and(words: list[str]) -> str:
    return words[0] if len(words) == 1 else ", ".join(words[:-1]) + " and " + words[-1]


def example_only(name: str) -> str:
    return f"{name} holds only the example row: add your rows and give it again."


def separate(name: str, which: str, generic_names: list[str], school_names: list[str]) -> str:
    """The card of a file in a drop that mixes generic and school files; `which` is "generic" or "school"."""
    return (f"{name} is one of the {which} CSV files; give the generic files ({', '.join(generic_names)}) and "
            f"the school files ({', '.join(school_names)}) in separate drops.")


def second(name: str, kind: str) -> str:
    return f"{name} is a second {kind}.csv in this drop: give one {kind}.csv at a time."


def sort_out(files: dict[str, tuple[str, bytes]], stored: dict) -> dict:
    """What a drop's generic files ({kind: (name, bytes)}, in drop order) come to, against the stored plan:
    {"refused": {kind: message}}, or {"name", "data", "read": [kinds], "left_out": [kinds]} — the reading's
    name and workbook, the kinds in it, and the files holding only example rows (their part of the plan
    stays as stored). Nothing but example rows anywhere is refused, so the stored plan is never replaced by
    an empty one; a staff, units or duties file missing (or example-only) with nothing stored for it is
    refused naming the files to add. A row with more cells than its header refuses the whole drop first:
    its cells would be lost, so no part of the drop is read."""
    over = {kind: rows for kind, (name, data) in files.items() if (rows := _overflowing(name, data))}
    if over:
        texts = [f"{files[k][0]} row {n}: {OVERFLOW}" for k, rows in over.items() for n in rows][:SHOWN_ROWS]
        why = "; ".join(texts)
        return {"refused": {kind: (f"Nothing in {name} could be read: {why}." if kind in over else
                                   f"{name} was not read: {why}. Give the files again.")
                            for kind, (name, _) in files.items()}}
    _, count = assemble(files)
    if count == 0:
        return {"refused": {kind: example_only(name) for kind, (name, _) in files.items()}}
    read = [kind for kind, (name, data) in files.items() if _real_rows(name, data)]
    missing = [kind for kind in REQUIRED if kind not in read and not (stored or {}).get(_SECTION[kind][0])]
    if missing:
        names = [f"{kind}.csv" for kind in missing]
        why = (f"the timetable holds no {' or '.join(_SECTION[k][1] for k in missing)} yet, so give "
               f"{_and(names)} with your rows in the same drop.")
        return {"refused": {kind: f"{name} was not read: {why}" for kind, (name, _) in files.items()}}
    used = {kind: files[kind] for kind in read}
    data, _ = assemble(used)
    return {"name": f"{SUMMARY} ({', '.join(name for name, _ in used.values())}).xlsx", "data": data,
            "read": read, "left_out": [kind for kind in files if kind not in read]}
