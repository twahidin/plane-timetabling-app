"""The plan written back out as a staff-deployment workbook, and the blank template (spec
docs/superpowers/specs/2026-09-26-deployment-board-design.md §3).

The importer is the contract: every header here is one `plan.importer` looks for, so importing an
export into an empty timetable gives the same plan back. What a workbook cannot hold (a split whose
parts are not one teacher each, a lesson longer than a quadruple, ...) is written as near as it can
be and listed on the README sheet, never dropped silently."""
from __future__ import annotations

import io
import re

from . import importer as I
from . import model as M

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXAMPLE_NOTE = "example: replace or delete"

_LESSON_COLS = (("Single", "1"), ("Double", "2"), ("Triple", "3"), ("Quadruple", "4"))
_STAFF_HEAD = ["Staff", "Short", "Dept", "Teaching Load Factor", "Allowance", "Reductions", "Provisional", "Shares"]
_BAD_TITLE = re.compile(r"[\[\]:*?/\\]")
# the importer's own sheet names, and "History", which Excel keeps for itself
_RESERVED_TITLES = I._SKIP_SHEETS | I._SPECIAL_SHEETS | {"history"}


# ---------------------------------------------------------------------------
# department sheets
# ---------------------------------------------------------------------------

def _text(value) -> str:
    return "" if value is None else str(value)


def _cell_name(req: dict) -> str:
    return " ".join(p for p in (_text(req.get("dept")), _text(req.get("subject")),
                                ",".join(_text(c) for c in req.get("classes") or ())) if p)


def _lesson_total(reqs: list[dict]) -> dict[str, int]:
    out: dict[str, int] = {}
    for r in reqs:
        for k, n in r["lessons"].items():
            if n:
                out[k] = out.get(k, 0) + n
    return out


def _nonzero(lessons: dict) -> dict:
    return {str(k): n for k, n in lessons.items() if n}


def _split_shares(parts: list[dict], base: str, lessons: dict) -> list[int] | None:
    """The `Split` shares that make the importer rebuild exactly these parts, or None when the
    column cannot say it: part ids `<base>-a`, `-b`, ... in order and the lessons shared out as
    `split_lessons` shares them. Any part's teachers can be written: one Teacher n column a part,
    "(unassigned)" or "A + B" (importer._share_teachers)."""
    ids = [r["id"] for r in parts]
    try:
        if ids != [M.split_part_id(base, i) for i in range(len(parts))]:
            return None
        shares = [r["periods"] for r in parts]
        if any(r["periods"] != sum(int(k) * n for k, n in r["lessons"].items()) for r in parts):
            return None
        rebuilt = M.split_lessons(lessons, shares)
    except (M.PlanError, TypeError, ValueError):
        return None
    if [_nonzero(p) for p in rebuilt] != [_nonzero(r["lessons"]) for r in parts]:
        return None
    return shares


def _own_base(reqs: list[dict]) -> tuple[str, bool]:
    """(the id the cell has as one requirement, whether it is subject-qualified): the plain
    requirement_id, or the subject-qualified id a second subject of the same department, level and
    classes has (model.requirement_id_for) — its own id, or its split parts' `<base>-a`, ..."""
    first = reqs[0]
    args = (first["dept"], _text(first["level"]), first["grouping"], [_text(c) for c in first["classes"]])
    plain = M.requirement_id(*args)
    qualified = M.requirement_id_for(*args, _text(first["subject"]), {plain})
    ids = [r["id"] for r in reqs]
    if ids == [qualified] or (len(ids) > 1 and ids[0] == M.split_part_id(qualified, 0)):
        return qualified, True
    return plain, False


def _cell_row(reqs: list[dict], names: dict[str, str], person: str, notes: list[str]) -> dict:
    """One workbook row for one board cell (plan.model.cell_key): a single requirement, a co-taught
    one, or every part of a split. Whether its id comes back is checked by write_workbook, which
    knows the rows before it."""
    reqs = sorted(reqs, key=lambda r: r["id"])
    first = reqs[0]
    name = _cell_name(first)
    base, qualified = _own_base(reqs)
    lessons = _lesson_total(reqs)
    teachers: list[str] = []
    for r in reqs:
        for t in r["teachers"]:
            if t not in teachers:
                teachers.append(t)
    for t in teachers:
        if t not in names:
            notes.append(f"{name}: {t} is not on the staff list, so it is written as its id")

    def written(tid: str) -> str:
        return names.get(tid) or tid

    split = None
    columns = [written(t) for t in teachers]
    if len(reqs) > 1:
        shares = _split_shares(reqs, base, lessons)
        if shares is not None:
            split = "/".join(str(s) for s in shares)
            columns = [" + ".join(written(t) for t in r["teachers"]) or I.UNASSIGNED for r in reqs]
            for t in (t for t in teachers if "+" in written(t)):
                notes.append(f"{name}: {written(t)} has a '+' in the name, which a Split row reads as two "
                             f"{person}s; rename them before uploading this workbook")
        else:
            notes.append(f"{name}: its parts ({', '.join(r['id'] for r in reqs)}) are not a split the Split "
                         f"column can say (part ids ending -a, -b, ... in order, lessons shared largest first), "
                         f"so they are written as one co-taught row")
        differ = [label for label, value in (("locked", lambda r: r["locked"]), ("size", lambda r: r.get("size")),
                                             ("venue kind", lambda r: r["venue"].get("kind")))
                  if len({value(r) for r in reqs}) > 1]
        if differ:
            notes.append(f"{name}: its parts differ in {', '.join(differ)}; the workbook has one value for the "
                         f"whole row, so they come back alike")

    long = sorted((int(k) for k in lessons if k not in dict(_LESSON_COLS).values()))
    if long:
        notes.append(f"{name}: lessons longer than 4 periods ({', '.join(str(k) for k in long)}) have no column; "
                     f"importing this workbook reports its total as not matching its lessons")

    grouping = first["grouping"]
    classes = [_text(c) for c in first["classes"]]
    return {
        "level": _text(first["level"]) or None,
        "subject": _text(first["subject"]) or None,
        "total": sum(int(r["periods"] or 0) for r in reqs),
        "lessons": lessons,
        "grouping": "Entire Class" if grouping == "class" and len(classes) == 1 else grouping,
        "classes": classes,
        "teachers": columns,
        "split": split,
        "size": next((r["size"] for r in reqs if r.get("size") is not None), None),
        "venue_kind": next((r["venue"]["kind"] for r in reqs if r["venue"].get("kind")), None),
        "locked": any(r["locked"] for r in reqs),
        "sort": (_text(first["level"]), _text(first["subject"]), _text(grouping), classes[0] if classes else ""),
        "note": None,
        "key": (_text(first["level"]), grouping, classes),      # with the sheet's title: the id it comes back as
        "name": name,
        "id": base if len(reqs) > 1 else first["id"],          # the id (or split base) it has in the plan
        # its id is the plain requirement_id: the importer gives that to the first row that asks
        "plain": not qualified and (base if len(reqs) > 1 else first["id"]) == base,
    }


def _import_order(rows: list[dict]) -> None:
    """Set each row's `sort` to the order the importer must read it in. Rows go by level, subject,
    grouping and class, except that a row of the same classes as the row with the plain id (a second
    subject, its id subject-qualified by model.requirement_id_for) follows that row: the importer
    gives the plain id to whichever comes first."""
    plain_rows = {}
    for r in rows:
        if not r["note"] and r["plain"]:
            level, grouping, classes = r["key"]
            plain_rows.setdefault((level, grouping, tuple(classes)), r["sort"])
    for r in rows:
        own = r["sort"]
        sibling = None
        if not r["note"] and not r["plain"]:
            level, grouping, classes = r["key"]
            sibling = plain_rows.get((level, grouping, tuple(classes)))
        r["sort"] = (own, 0, ()) if sibling is None else (sibling, 1, own)


def _sheet_title(dept: str, used: set[str], notes: list[str]) -> str:
    """The department's own name as the sheet's title (the importer reads the title as the
    department), unless Excel or the importer cannot take it."""
    title = _BAD_TITLE.sub("-", dept).strip().strip("'").strip()[:31] or "Department"
    low = title.lower()
    if low in _RESERVED_TITLES:
        title = f"{title} dept"[:31]
    elif low.startswith("sheet"):           # the importer skips Excel's "Sheet1", ...: lead with a word
        title = f"Dept {title}"[:31]
    base, n = title, 2
    while title.lower() in used:
        title = f"{base[:28]} {n}"
        n += 1
    used.add(title.lower())
    if title != dept:
        notes.append(f"department {dept!r} is written on sheet {title!r} and comes back under that name")
    return title


def _write_department(ws, rows: list[dict], *, min_classes: int = 1, min_teachers: int = 1) -> None:
    noted = any(r["note"] for r in rows)
    n_classes = max([min_classes, *(len(r["classes"]) for r in rows)])
    n_teachers = max([min_teachers, *(len(r["teachers"]) for r in rows)])
    head = ["Level", "Subject", "Total", *(label for label, _ in _LESSON_COLS), "Grouping",
            *(f"Class {i}" for i in range(1, n_classes + 1)), *(f"Teacher {i}" for i in range(1, n_teachers + 1)),
            "Split", "Total Students", "Venue kind", "Locked"]
    if noted:
        head.append("Note")
    ws.append(head)
    for r in sorted(rows, key=lambda r: r["sort"]):
        ws.append([
            r["level"], r["subject"], r["total"], *((r["lessons"].get(k) or None) for _, k in _LESSON_COLS),
            r["grouping"],
            *(r["classes"] + [None] * (n_classes - len(r["classes"]))),
            *(r["teachers"] + [None] * (n_teachers - len(r["teachers"]))),
            r["split"], r["size"], r["venue_kind"], "yes" if r["locked"] else None,
            *([r["note"]] if noted else []),
        ])


# ---------------------------------------------------------------------------
# Groups sheet
# ---------------------------------------------------------------------------

def _division_codes(plan: dict, notes: list[str]) -> dict[str, list[str]]:
    """Division id -> the group codes written on its row. The importer makes one band a code family
    ("1ELP", "1ELR" -> EL) on each row, taking every requirement whose grouping is one of the codes;
    a band it would rebuild differently is named on the README."""
    reqs = {r["id"]: r for r in plan["requirements"]}
    division_ids = {d["id"] for d in plan["divisions"]}
    out: dict[str, list[str]] = {}
    families_seen: dict[str, set[str]] = {}
    for band in plan["bands"]:
        div = band["division"]
        if div not in division_ids:
            notes.append(f"band {band['id']}: its division {div!r} is not in the plan, so it is left out")
            continue
        codes: list[str] = []
        for opt in band["options"]:
            grouping = (reqs.get(opt) or {}).get("grouping")
            if grouping and grouping != "class" and grouping not in codes:
                codes.append(grouping)
        families = {I.family(c) for c in codes}
        takes = {r["id"] for r in plan["requirements"] if r["grouping"] in codes}
        exact = (len(families) == 1 and band["id"] == M.band_id(div, next(iter(families)))
                 and takes == set(band["options"]) and not (families & families_seen.get(div, set())))
        if not exact:
            notes.append(f"band {band['id']}: the Groups sheet cannot hold it exactly (its options' group codes "
                         f"{', '.join(codes) or 'none'}); check it after importing this workbook")
        families_seen.setdefault(div, set()).update(families)
        row = out.setdefault(div, [])
        row.extend(c for c in codes if c not in row)
    return out


def _write_groups(ws, plan: dict, notes: list[str]) -> None:
    codes = _division_codes(plan, notes)
    divisions = plan["divisions"]
    n_classes = max([1, *(len(d["classes"]) for d in divisions)])
    n_codes = max([1, *(len(c) for c in codes.values())])
    ws.append(["Class(es) in a Division", *([None] * (n_classes - 1)), "Groups (Subjects)", *([None] * (n_codes - 1))])
    ws.append([*range(1, n_classes + 1), *(f"Group {i}" for i in range(1, n_codes + 1))])
    for d in divisions:
        classes = [_text(c) for c in d["classes"]]
        if d["id"] != M.division_id(classes):
            notes.append(f"division {d['id']} comes back as {M.division_id(classes)} when this workbook is imported")
        row_codes = codes.get(d["id"], [])
        ws.append([*classes, *([None] * (n_classes - len(classes))), *row_codes])


# ---------------------------------------------------------------------------
# Control sheet
# ---------------------------------------------------------------------------

def _reductions_text(reductions: list[dict]) -> str:
    return "; ".join(f"{r['reason']} {r['periods']}".strip() for r in reductions)


def _shares_text(shares: dict) -> str:
    return "; ".join(f"{d} {n}" for d, n in sorted((shares or {}).items()))


def _write_control(ws, staff: list[dict], notes: list[str], example: bool = False) -> None:
    ws.append(_STAFF_HEAD + (["Note"] if example else []))
    rows = sorted(staff, key=lambda s: (_text(s["dept"]).lower(), _text(s["name"]).lower(), s["id"]))
    by_name: dict[str, list[str]] = {}
    for s in rows:
        name = _text(s["name"]).strip() or s["id"]
        by_name.setdefault(name.lower(), []).append(name)
    for same in by_name.values():
        if len(same) > 1:
            count = {2: "two", 3: "three"}.get(len(same), str(len(same)))
            notes.append(f"{count} staff are named {same[0]}: their rows may merge on upload — give them "
                         f"different names")
    for s in rows + (_EXAMPLE_STAFF if example else []):
        name = _text(s["name"]).strip() or s["id"]
        if I.is_unassigned(name):
            notes.append(f"{name}: a Teacher column of a Split row reads this name as a share nobody teaches; "
                         f"rename them before uploading this workbook")
        text = _reductions_text(s["reductions"])
        if I._read_reductions(text, name, []) != s["reductions"]:
            notes.append(f"{name}: reductions {text!r} do not read back as written (a reason holding ';', ',' "
                         f"or a number of its own); check them after importing this workbook")
        shares = _shares_text(s.get("shares"))
        if I._read_shares(shares, name, []) != dict(sorted((s.get("shares") or {}).items())):
            notes.append(f"{name}: shares {shares!r} do not read back as written (a department name holding ';' "
                         f"or ','); check them after importing this workbook")
        ws.append([name, _text(s["short"]) or None, _text(s["dept"]) or None, s["load_factor"], s["allowance"],
                   text or None, "yes" if s["provisional"] else None, shares or None,
                   *([s.get("note")] if example else [])])


# ---------------------------------------------------------------------------
# README
# ---------------------------------------------------------------------------

def _readme_lines(plan: dict, settings: dict, notes: list[str], example: bool) -> list[str]:
    v = plan.get("vocabulary") or M.DEFAULT_VOCABULARY
    person, group, lesson = v["person"], v["group"], v["requirement"]
    rules = (settings or {}).get("rules") or {}
    try:
        capacity = f"{rules['max_load']} × {M.cycle_days(settings)} periods"
    except (KeyError, TypeError, ZeroDivisionError):
        capacity = "the most periods a day × the days in the cycle"
    lines = [
        "Staff deployment workbook",
        "",
        "Fill this in and upload it on the Plan tab (Upload workbook). Uploading it again later merges: "
        "rows marked Locked are never changed by an upload.",
        "",
        f"Department sheets (one per department, the sheet's name is the department): one row per {lesson}.",
        "  Level: the level or stream, e.g. 1G3.  Subject: the subject's name.",
        "  Total: periods a cycle.  Single, Double, Triple, Quadruple: how many lessons of 1, 2, 3 and 4 periods; "
        "they must add up to Total.",
        f"  Grouping: Entire Class (one {lesson} for each {group} named), or a group code such as 1ELP for a "
        f"{lesson} drawn from more than one {group}.",
        f"  Class 1, Class 2, ...: one {group} a column.  Teacher 1, Teacher 2, ...: one {person} a column, by name.",
        f"  Split: how the periods are shared out when each {person} takes their own part, one share a Teacher "
        f"column in column order, e.g. 3/2. With a Split, a Teacher column may read {I.UNASSIGNED} (a share "
        f"nobody teaches yet) or 'Kelly Wong + Eugene Lee' (a share two {person}s teach together). Leave Split "
        f"empty when everyone named teaches every lesson together.",
        "  Total Students: the group's size.  Venue kind: the kind of room it needs, e.g. lab.",
        f"  Locked: yes when the {lesson} is signed off; an upload will not change it.",
        "",
        "Groups: one row per division. Class(es) in a Division lists the classes whose option groups run "
        "at the same time; Groups (Subjects) lists those groups' codes. Codes of one family (1ELP, 1ELR) "
        "make one band.",
        "",
        f"Control: one row per {person}. Teaching Load Factor: 1 for full time, 0.5 for half. "
        f"Allowance: periods a cycle before reductions (empty: load factor × {capacity}). "
        "Reductions: each role and its periods, separated by ';', e.g. HOD 8; CCA 2. "
        f"Provisional: yes for a placeholder {person} still to be named. "
        f"Shares: periods a cycle the {person} gives other departments, e.g. MATH 10; SCI 6 (a department "
        f"named here may assign them).",
    ]
    if example:
        lines += ["", f"Rows whose Note says '{EXAMPLE_NOTE}' show the shape: replace them with your own or delete them."]
    if notes:
        lines += ["", "Not written exactly (check these after importing this workbook):", *(f"  {n}" for n in notes)]
    return lines


# ---------------------------------------------------------------------------
# write_workbook
# ---------------------------------------------------------------------------

def _example_rows() -> list[dict]:
    common = {"level": "1", "subject": "Mathematics", "total": 5, "lessons": {"1": 1, "2": 2},
              "grouping": "Entire Class", "venue_kind": None, "locked": False}
    return [
        {**common, "classes": ["1A", "1B"], "teachers": ["Example Teacher A"], "split": None, "size": 36,
         "sort": ("1", "Mathematics", "class", "1A"), "note": EXAMPLE_NOTE},
        {**common, "classes": ["1C"], "teachers": ["Example Teacher A", "New MATH teacher 1"], "split": "3/2",
         "size": 35, "sort": ("1", "Mathematics", "class", "1C"), "note": EXAMPLE_NOTE},
    ]


_EXAMPLE_STAFF = [
    {"id": "example-teacher-a", "name": "Example Teacher A", "short": "ETA", "dept": "MATH", "load_factor": 1.0,
     "allowance": 30, "reductions": [{"reason": "HOD", "periods": 8}], "provisional": False, "note": EXAMPLE_NOTE},
    {"id": "new-math-teacher-1", "name": "New MATH teacher 1", "short": "", "dept": "MATH", "load_factor": 1.0,
     "allowance": None, "reductions": [], "provisional": True, "note": EXAMPLE_NOTE},
]


def write_workbook(plan: dict, settings: dict, *, example: bool = False) -> bytes:
    """The plan as the importer's own workbook: README, one sheet per department (sorted), Groups
    and Control. `example` adds the template's example MATH sheet and Control rows. PlanError when
    the plan does not normalise."""
    import openpyxl

    plan = M.normalise(plan)
    notes: list[str] = []
    names = {s["id"]: _text(s["name"]).strip() or s["id"] for s in plan["staff"]}

    cells: dict[tuple, list[dict]] = {}
    for r in plan["requirements"]:
        if not r["classes"]:
            notes.append(f"{_cell_name(r) or r['id']}: names no {plan['vocabulary']['group']}, so it is left out")
            continue
        cells.setdefault(M.cell_key(r), []).append(r)
    by_dept: dict[str, list[dict]] = {}
    for key, reqs in cells.items():
        by_dept.setdefault(_text(key[0]), []).append(_cell_row(reqs, names, plan["vocabulary"]["person"], notes))
    if example:
        by_dept.setdefault("MATH", []).extend(_example_rows())

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("README")
    used: set[str] = {"readme", "groups", "control"}
    seen: dict[str, str] = {}
    owners: dict[str, str] = {}            # the importer's own record of the ids it has given out
    for dept in sorted(by_dept, key=str.lower):
        title = _sheet_title(dept, used, notes)
        rows = by_dept[dept]
        _import_order(rows)
        for row in sorted(rows, key=lambda r: r["sort"]):
            if row["note"]:                     # the template's example rows
                continue
            level, grouping, classes = row["key"]
            # the id the importer gives this row, reading the sheet's title as its department
            req_id = I.import_id(title, level, grouping, classes, _text(row["subject"]).strip(), owners)
            what = f"{title} {row['subject'] or ''} {','.join(classes)}".replace("  ", " ")
            if req_id in seen:
                notes.append(f"{seen[req_id]} and {what} both come back as {req_id} when this workbook is "
                             f"imported: one replaces the other")
                continue
            seen[req_id] = what
            if title == dept and req_id != row["id"]:
                notes.append(f"{row['name']}: {row['id']} comes back as {req_id} when this workbook is imported")
        room = {"min_classes": 4, "min_teachers": 3} if example else {}     # room to fill the template in
        _write_department(wb.create_sheet(title), rows, **room)
    _write_groups(wb.create_sheet("Groups"), plan, notes)
    _write_control(wb.create_sheet("Control"), plan["staff"], notes, example)
    for line in _readme_lines(plan, settings, notes, example):
        readme.append([line or None])
    readme.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
