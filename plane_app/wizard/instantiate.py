"""Turn a chosen template and its knobs into everything the wizard hands out (spec
docs/superpowers/specs/2026-09-20-start-wizard-design.md §4-§5).

The model chooses a template and sets its knobs; `library.check_knobs` validates them. From there
everything is deterministic and rendered here: the timetable's settings (slot labels for the whole
cycle, the rules in slots), the facts the side panel shows, a small sample plan and the draft
organisation it generates into, the workbook the user fills in (readable by `plan.importer`), a
one-page printable guide and a sample PDF of the printed timetable.

Two shapes of sample, following the template's `sheets`:

* `deployment` — a school: classes in levels, one requirement per subject per class, and option
  groups (two sets by default) where a subject is banded.
* `generic` — a ward, clinic, office, court or tournament: the sample's units become classes, and
  every occurrence of a duty becomes its own requirement with its own team of staff, so the load
  spreads across the whole sample roster rather than pinning one team to every shift.

A plan lesson runs one to four slots (`plan.model`), so a duty longer than that is carried as
consecutive blocks — a twelve-hour shift is three blocks of four hours. The sample exists to show
what a printed timetable looks like, and the engine's greedy build must place it clean: a sample
that does not is a library bug, and `sample_pdf` raises `WizardError` rather than print a hole.
"""
from __future__ import annotations

import copy
import io
from datetime import date as _date

from .. import intake
from ..plan import model as M
from ..plan.generate import generate
from ..print import grid as G, pdf as P
from .library import WizardError

_LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")

MAX_LESSON = 4          # plan.model's lesson lengths: a longer duty is carried as several blocks
MAX_SAMPLE_GRIDS = 6    # the sample PDF is a taste, not a timetable
MAX_SAMPLE_ROWS = 3     # sample rows a sheet, and never the same row twice
CLASS_SIZE = 30         # invented, like every name in a sample
ROOM_CAP = 40
SPARE_ROOMS = 2         # rooms beyond one per class, so a band's options always have somewhere to go
SHARE_ABOVE = 0.6       # a venue kind this full of its own sample is marked shared: duties overlap in it

# Settings keys the wizard leaves exactly as the timetable already has them.
CARRIED_SETTINGS = ("calendar", "solve", "provider", "engine")


# ---------------------------------------------------------------------------
# time: the shape of the cycle, and its labels
# ---------------------------------------------------------------------------

def _minutes(hhmm: str) -> int:
    hours, minutes = str(hhmm).split(":")
    return int(hours) * 60 + int(minutes)


def _clock(start: int, index: int, slot_minutes: int) -> str:
    total = (start + index * slot_minutes) % (24 * 60)
    return f"{total // 60:02d}:{total % 60:02d}"


def shape(template: dict, knobs: dict) -> dict:
    """The cycle the knobs describe: slot length, slots a day, when the day starts, days in the
    cycle. The template's `time` is the default; a knob that names one of them wins."""
    time = template["time"]
    slot_minutes = knobs.get("slot_minutes") or knobs.get("period_minutes") or time["slot_minutes"]
    day_start = knobs.get("open_from") or time["day_start"]
    if "periods_per_day" in knobs:
        slots_per_day = int(knobs["periods_per_day"])
    elif "open_to" in knobs:
        open_minutes = _minutes(knobs["open_to"]) - _minutes(day_start)
        slots_per_day = max(1, open_minutes // int(slot_minutes))
    else:
        slots_per_day = time["slots_per_day"]
    return {"slot_minutes": int(slot_minutes), "slots_per_day": int(slots_per_day),
            "day_start": day_start, "cycle_days": int(knobs["cycle_days"]),
            "days_per_week": time["days_per_week"], "kind": time["labels"]}


def _unit_word(template: dict, knobs: dict) -> str:
    """The word for a slot in anything the user reads: an hour-long slot is an hour, anything else a
    period. "Slot" is the engine's word and never appears in the workbook, the guide or the panel."""
    return "hours" if shape(template, knobs)["slot_minutes"] == 60 else "periods"


def _day_name(sh: dict, day: int) -> str:
    """"Mon"; "Odd Mon"/"Even Mon" over a fortnight; "W3 Mon" beyond; "Day 3" with no weekdays."""
    cycle_days, per_week = sh["cycle_days"], sh["days_per_week"]
    if sh["kind"] == "period":
        return "" if cycle_days == 1 else f"Day {day + 1}"
    name = WEEKDAYS[day % per_week]
    if cycle_days <= per_week:
        return name
    week = day // per_week
    weeks = -(-cycle_days // per_week)
    return f"{('Odd', 'Even')[week] if weeks == 2 else f'W{week + 1}'} {name}"


def labels_for(template: dict, knobs: dict) -> list[str]:
    """One label a slot for the whole cycle: "Odd Mon P1"…, "Mon 07:00"…, "P1"…"""
    sh = shape(template, knobs)
    start = _minutes(sh["day_start"])
    out = []
    for day in range(sh["cycle_days"]):
        prefix = _day_name(sh, day)
        for i in range(sh["slots_per_day"]):
            slot = _clock(start, i, sh["slot_minutes"]) if sh["kind"] == "day-hour" else f"P{i + 1}"
            out.append(f"{prefix} {slot}" if prefix else slot)
    return out


def rules_for(template: dict, knobs: dict) -> dict:
    """The template's slot rules, in the settings' words, with `slots_per_day` so that load, runs
    and breaks are judged within a day rather than across the whole cycle."""
    sh = shape(template, knobs)
    slots_per_day = sh["slots_per_day"]
    rules = template["rules"]
    max_run = rules["max_run_slots"]
    if "shift_hours" in knobs:
        max_run = max(1, round(int(knobs["shift_hours"]) * 60 / sh["slot_minutes"]))
    rest = [int(knobs["recess_after"])] if "recess_after" in knobs else list(rules["mandatory_rest"])
    return {"max_load": min(rules["max_load_slots"], slots_per_day),
            "max_run": min(max_run, slots_per_day),
            "mandatory_rest": [t for t in rest if 0 <= t < slots_per_day],
            "slots_per_day": slots_per_day}


def settings_for(template: dict, knobs: dict, current_settings: dict) -> dict:
    """A full settings document for the timetable: time and rules from the template and knobs,
    everything else (the calendar, the solver, the provider, the engine) as it already is."""
    sh = shape(template, knobs)
    settings = {
        "time": {"slot_minutes": sh["slot_minutes"], "slots_per_day": sh["slots_per_day"],
                 "start": sh["day_start"], "days_per_week": sh["days_per_week"],
                 "labels": labels_for(template, knobs), "window": "day"},
        "rules": rules_for(template, knobs),
    }
    for key in CARRIED_SETTINGS:
        if key in (current_settings or {}):
            settings[key] = copy.deepcopy(current_settings[key])
    return settings


# ---------------------------------------------------------------------------
# the sample plan
# ---------------------------------------------------------------------------

def _people(word: str, count: int) -> list[str]:
    """"Nurse A", "Nurse B", … — invented names, never anyone real."""
    return [f"{word.title()} {_LETTERS[i % len(_LETTERS)]}" for i in range(count)]


def _initials(name: str) -> str:
    letters = [w[0] for w in str(name).split() if w and w[0].isalpha()]
    return ("".join(letters) or "X").upper()


def _staff(names: list[str]) -> list[dict]:
    return [{"id": M.slug(name), "name": name, "short": _initials(name), "dept": "",
             "load_factor": 1.0, "avail": None, "max_periods_day": None, "source_hash": None}
            for name in names]


def _chunks(length: int) -> list[int]:
    """A duty of `length` slots as lessons of at most four, as evenly as they divide: 12 -> 4+4+4,
    6 -> 3+3, 5 -> 3+2."""
    parts = max(1, -(-int(length) // MAX_LESSON))
    base, extra = divmod(int(length), parts)
    return [base + 1] * extra + [base] * (parts - extra)


def _lessons(lengths: list[int]) -> dict:
    """Lesson lengths counted the way the plan holds them: [3, 3] -> {"3": 2}."""
    out: dict[str, int] = {}
    for length in lengths:
        out[str(length)] = out.get(str(length), 0) + 1
    return out


def _requirement(dept: str, level: str, subject: str, grouping: str, classes: list[str],
                 teachers: list[str], lessons: dict, kind: str | None, size: int | None) -> dict:
    return {"id": M.requirement_id(dept, level, grouping, classes), "dept": dept, "level": level,
            "subject": subject, "periods": sum(int(k) * v for k, v in lessons.items()),
            "lessons": dict(lessons), "grouping": grouping, "classes": list(classes),
            "teachers": list(teachers), "size": size,
            "venue": {"kind": kind, "room": None}, "sync_with": None, "source_hash": None}


class _Roster:
    """Hands out teams of staff round the sample's roster, so every person carries a share."""

    def __init__(self, ids: list[str]):
        self.ids, self.at = ids, 0

    def team(self, size: int) -> list[str]:
        size = max(1, min(int(size), len(self.ids)))
        team = [self.ids[(self.at + i) % len(self.ids)] for i in range(size)]
        self.at = (self.at + size) % len(self.ids)
        return team


def _subject_code(name: str) -> str:
    """A short letters-only code for a subject: its initials when it has several words ("Module A
    tutorial" -> MAT), else its first three letters ("Art" -> ART). Option groups are named
    <level><code><letter> ("1MATA"), which is how the deployment importer reads a band: the letters
    between the level and the last character are the family the options share."""
    words = [w for w in str(name).split() if w]
    code = "".join(w[0] for w in words) if len(words) > 1 else str(name)[:3]
    return "".join(ch for ch in code.upper() if ch.isalnum()) or "X"


def _subject_codes(subjects: list[dict]) -> dict[str, str]:
    """One code a subject, distinct within the template: two subjects that share a code would share
    an option group, and their bands would collide."""
    out, seen = {}, set()
    for subject in subjects:
        code = _subject_code(subject["name"])
        while code in seen:
            code += "X"
        seen.add(code)
        out[subject["name"]] = code
    return out


def _split_parts(periods) -> tuple[int, int]:
    periods = int(periods)
    return periods - periods // 2, periods // 2


def _split_of(subject: dict) -> str:
    """The Split column's example: the subject's periods shared between two teachers, "3/2"."""
    return "%d/%d" % _split_parts(subject["periods"])


def _longest_run(rules: dict) -> int:
    """The longest lesson a day of this cycle can actually hold: the widest stretch between its
    mandatory breaks, and no longer than the run and load a plane is allowed. A `Split` puts each
    teacher's share into one lesson of that many slots (`plan.importer._apply_split`), so a sample
    Split longer than this would hand the user a workbook that cannot be placed."""
    breaks = sorted({t for t in rules["mandatory_rest"] if 0 <= t < rules["slots_per_day"]})
    longest, start = 0, 0
    for t in (*breaks, rules["slots_per_day"]):
        longest = max(longest, t - start)
        start = t + 1
    return max(1, min(longest, rules["max_run"], rules["max_load"]))


def _options_per_band(knobs: dict) -> int:
    for key in ("sets_per_subject", "tutorial_groups_per_module"):
        if key in knobs:
            return max(1, int(knobs[key]))
    return 2


def education_sample(template: dict, knobs: dict) -> dict:
    """levels x classes, one requirement a subject a class, option groups where a subject is banded."""
    sample, vocabulary = template["sample"], template["vocabulary"]
    levels = [(f"Year {n + 1}", [f"{n + 1}{_LETTERS[c]}" for c in range(sample["classes_per_level"])])
              for n in range(sample["levels"])]
    staff = _staff(_people(vocabulary["person"], sample["teachers"]))
    roster = _Roster([s["id"] for s in staff])
    banded = knobs.get("bands", True)
    options = _options_per_band(knobs)
    subject_codes = _subject_codes(sample["subjects"])

    classes, divisions, requirements, bands = [], [], [], []
    for index, (level, class_codes) in enumerate(levels, start=1):
        classes += [{"code": code, "level": level, "size": CLASS_SIZE} for code in class_codes]
        division = {"id": M.division_id(class_codes), "classes": list(class_codes)}
        divisions.append(division)
        for subject in sample["subjects"]:
            lengths = {k: v for k, v in subject["lengths"].items() if v}
            if subject["band"] and banded:
                option_ids = []
                for option in range(options):
                    group_code = f"{index}{subject_codes[subject['name']]}{_LETTERS[option]}"
                    requirement = _requirement(subject["name"], level, subject["name"], group_code,
                                               class_codes, roster.team(1), lengths, "classroom", CLASS_SIZE)
                    requirements.append(requirement)
                    option_ids.append(requirement["id"])
                bands.append({"id": M.band_id(division["id"], subject_codes[subject["name"]]),
                              "division": division["id"], "options": option_ids})
                continue
            for code in class_codes:
                requirements.append(_requirement(subject["name"], level, subject["name"], "class",
                                                 [code], roster.team(1), lengths, "classroom", CLASS_SIZE))
    return {"version": M.PLAN_VERSION, "staff": staff, "classes": classes, "divisions": divisions,
            "requirements": requirements, "bands": bands,
            "rules": {"edge_subjects": [], "no_double_across_rest": True, "pinned": []},
            "source": {"file": f"{template['name']} sample", "imported": None}}


def generic_sample(template: dict, knobs: dict) -> dict:
    """Units become classes; every occurrence of a duty is its own requirement with its own team,
    so a fortnight of shifts spreads over the roster instead of landing on one crew.

    The sample's size comes from the template, not from the knobs that count rooms, desks or squads:
    the library validates that size (and that every duty fits inside a day), so the sample stays one
    the engine can place whatever the user answers."""
    sample, vocabulary = template["sample"], template["vocabulary"]
    units = [f"{vocabulary['group'].title()} {i + 1}" for i in range(sample["units"])]
    staff = _staff(_people(vocabulary["person"], sample["staff"]))
    roster = _Roster([s["id"] for s in staff])

    classes = [{"code": unit, "level": "", "size": None} for unit in units]
    divisions = [{"id": M.division_id([unit]), "classes": [unit]} for unit in units]
    requirements = []
    for unit in units:
        for duty in sample["duties"]:
            lengths = _chunks(duty["length_slots"])
            for occurrence in range(duty["per_cycle"]):
                code = f"{_initials(duty['name'])}{occurrence + 1}"
                requirements.append(_requirement(
                    duty["name"], unit, duty["name"], code, [unit], roster.team(duty["min_staff"]),
                    _lessons(lengths), duty["venue_kind"], None))
    return {"version": M.PLAN_VERSION, "staff": staff, "classes": classes, "divisions": divisions,
            "requirements": requirements, "bands": [],
            "rules": {"edge_subjects": [], "no_double_across_rest": True, "pinned": []},
            "source": {"file": f"{template['name']} sample", "imported": None}}


def sample_plan(template: dict, knobs: dict) -> dict:
    """A normalised plan document the sample organisation is generated from."""
    builder = education_sample if template["sheets"] == "deployment" else generic_sample
    return M.normalise(builder(template, knobs))


# ---------------------------------------------------------------------------
# the sample organisation
# ---------------------------------------------------------------------------

def _usable_slots(template: dict, knobs: dict) -> int:
    rules = rules_for(template, knobs)
    return max(1, rules["slots_per_day"] - len(rules["mandatory_rest"]))


def sample_venues(template: dict, knobs: dict, plan: dict) -> list[dict]:
    """The sample's rooms. Education invents one a class plus a couple of spares (a band's options
    sit in different rooms at the same time); the generic templates take the ones the sample names.
    A kind whose sample fills most of it is marked shared: several duties do overlap in one ward."""
    sh = shape(template, knobs)
    if template["sheets"] == "deployment":
        count = len(plan["classes"]) + SPARE_ROOMS
        return [{"id": f"room-{i + 1}", "name": f"Room {i + 1}", "cap": ROOM_CAP,
                 "kind": "classroom", "shared": False, "rest": False} for i in range(count)]

    venues = template["sample"]["venues"]
    per_kind: dict[str, int] = {}
    for venue in venues:
        per_kind[venue["kind"]] = per_kind.get(venue["kind"], 0) + 1
    demand: dict[str, int] = {}
    for r in plan["requirements"]:
        kind = r["venue"].get("kind")
        if kind:
            demand[kind] = demand.get(kind, 0) + r["periods"]
    supply = {kind: count * sh["cycle_days"] * _usable_slots(template, knobs)
              for kind, count in per_kind.items()}
    return [{"id": M.slug(v["name"]), "name": v["name"], "cap": v["capacity"], "kind": v["kind"],
             "shared": demand.get(v["kind"], 0) > SHARE_ABOVE * supply[v["kind"]], "rest": False}
            for v in venues]


def organisation_for(template: dict, knobs: dict, settings: dict, plan: dict) -> dict:
    """The draft organisation a plan generates into under this template, rooms and all: the sample's
    own plan, or the one the importer reads back out of the workbook the user was handed."""
    base = intake.empty_organisation(settings, f"{template['name']} sample")
    base["locations"] = sample_venues(template, knobs, plan) + list(base["locations"])
    org, _summary = generate(plan, base, settings)
    return org


def sample_organisation(template: dict, knobs: dict, settings: dict) -> dict:
    """The draft organisation the sample plan generates into, rooms and all."""
    return organisation_for(template, knobs, settings, sample_plan(template, knobs))


# ---------------------------------------------------------------------------
# facts
# ---------------------------------------------------------------------------

def facts(template: dict, knobs: dict) -> dict:
    """What the side panel shows about the configuration being considered."""
    sh = shape(template, knobs)
    plan = sample_plan(template, knobs)
    groups = len(plan["classes"]) + len({(r["grouping"], tuple(r["classes"]))
                                         for r in plan["requirements"] if r["grouping"] != "class"})
    return {
        "cycle_days": sh["cycle_days"],
        "slots_per_day": sh["slots_per_day"],
        "slots": sh["slots_per_day"] * sh["cycle_days"],
        "slot_minutes": sh["slot_minutes"],
        "unit": _unit_word(template, knobs),     # "hours" or "periods": never the engine's "slots"
        "sample": {"people": len(plan["staff"]), "groups": groups,
                   "requirements": len(plan["requirements"]),
                   "events": sum(sum(r["lessons"].values()) for r in plan["requirements"])},
        "tradeoffs": list(template["tradeoffs"]),
        "vocabulary": dict(template["vocabulary"]),
        "questions": list(template["questions"]),
    }


# ---------------------------------------------------------------------------
# the workbook and the guide: one description of the sheets, rendered twice
# ---------------------------------------------------------------------------

def _column(name, what: str, example="") -> dict:
    return {"name": name, "what": what, "example": example}


def _deployment_columns(v: dict) -> list[dict]:
    columns = [
        _column("Subject Level", f"The year or level of the {v['group']}es in this row", "Year 1"),
        _column("Subject", f"The {v['requirement']}, with its periods a cycle in brackets", "English (5 Periods)"),
        _column("Single", "How many lessons of one period", 3),
        _column("Double", "How many lessons of two periods together", 1),
        _column("Triple", "How many lessons of three periods together", 0),
        _column("Quadruple", "How many lessons of four periods together", 0),
        _column("Total", "Periods a cycle: the lengths above must add up to it", 5),
        _column("Grouping", f"\"Entire Class\", or the code of an option group when the {v['group']} splits", "Entire Class"),
    ]
    columns.append(_column("Class 1", f"The {v['group']} taught in this row", "1A"))
    columns += [_column(f"Class {n}", f"A further {v['group']} taught together with the first", "")
                for n in range(2, 7)]
    columns += [
        _column("Total Students", "How many students the row teaches", CLASS_SIZE),
        _column("Teacher 1", f"The {v['person']} who takes it", "Teacher A"),
        _column("Teacher 2", f"A second {v['person']} when two of them take it together", "Teacher C"),
        _column("Split", f"\"3/2\" splits the periods between the two {v['person']}s; leave it empty "
                         f"when they teach every lesson together", "3/2"),
    ]
    return columns


def _generic_sheets(template: dict, knobs: dict) -> list[dict]:
    """Sample rows for the four generic sheets: at most three a sheet (the Staff sheet lists whoever
    the duties need), every one of them distinct, and the four consistent with each other — the duties
    ask only for the venue kinds listed here, sit on the units listed here, and are carried by the
    people listed here. The rows are a rota that actually builds, so dropping the workbook straight
    back closes the loop the spec promises: workbook -> importer -> generate -> build, clean.
    `Co-staffed` never exceeds the row's own "Who can do it" pool."""
    v = template["vocabulary"]
    sample = template["sample"]
    sh, rules = shape(template, knobs), rules_for(template, knobs)

    venues = list(sample["venues"])[:MAX_SAMPLE_ROWS]
    kinds = {venue["kind"] for venue in venues}
    duties = ([d for d in sample["duties"] if d["venue_kind"] in kinds]
               or list(sample["duties"]))[:MAX_SAMPLE_ROWS]
    units = [f"{v['group'].title()} {i + 1}" for i in range(min(MAX_SAMPLE_ROWS, sample["units"]))] or \
            [f"{v['group'].title()} 1"]

    # A duty row is one requirement for its whole unit, and a unit is a plane like any other: the rows
    # sharing a unit cannot together ask for more than its cycle holds (cycle_days x max_load). The
    # sample plan gives every occurrence a group of its own and so never meets that ceiling; a filled-in
    # workbook does, which is why the sample rows count their occurrences down to what fits.
    on_unit = [units[i % len(units)] for i in range(len(duties))]
    rows_on = {unit: on_unit.count(unit) for unit in on_unit}
    days_for = {unit: max(1, sh["cycle_days"] // count) for unit, count in rows_on.items()}
    per_cycle = [max(1, min(d["per_cycle"],
                            days_for[unit] * max(1, rules["max_load"] // d["length_slots"])))
                 for d, unit in zip(duties, on_unit)]

    # One person a seat: the teams never overlap, so nobody carries two rows and each row's team is
    # exactly the pool its "Who can do it" lists.
    working = _people(v["person"], max(MAX_SAMPLE_ROWS, sum(d["min_staff"] for d in duties)))
    roster = _Roster(working)
    people = _people(v["person"], len(working) + 1)        # one more: the part-time, some-days row

    duty_rows = []
    for duty, unit, times in zip(duties, on_unit, per_cycle):
        team = roster.team(duty["min_staff"])
        # "Together with" is left empty: two rows of one unit starting together would ask that unit
        # for both at once, which is exactly what its own plane forbids. The README explains the column.
        duty_rows.append([duty["name"], unit, times, duty["length_slots"],
                          duty["venue_kind"], ", ".join(team), "",
                          len(team) if len(team) > 1 else "no"])

    spare = people[-1]                                     # on no duty row, so the example costs nothing
    staff_rows = [[name, f"Senior {v['person']}" if name == spare else v["person"].title(),
                   f"{WEEKDAYS[0]}-{WEEKDAYS[2]}" if name == spare else "All week",
                   0.5 if name == spare else 1.0]
                  for name in people]
    venue_rows = [[venue["name"], venue["kind"], venue["capacity"]] for venue in venues]
    unit_rows = [[unit, CLASS_SIZE] for unit in units]

    return [
        {"name": "Staff", "purpose": f"One row per {v['person']}.",
         "columns": [_column("Name", f"The {v['person']}'s name, as it should appear on the timetable", people[0]),
                     _column("Role", "What they do — used to group them on printouts", v["person"].title()),
                     _column("Availability", "When they work: \"All week\", or the days they are in", "Mon-Wed"),
                     _column("Load factor", "1 for full time, 0.5 for half, and so on", 1.0)],
         "rows": staff_rows},
        {"name": "Venues", "purpose": f"One row per {v['venue']} or other place a {v['requirement']} happens in.",
         "columns": [_column("Name", f"What the {v['venue']} is called", venues[0]["name"]),
                     _column("Kind", f"The sort of {v['venue']} it is; duties ask for a kind, not a name", venues[0]["kind"]),
                     _column("Capacity", "How many people fit", venues[0]["capacity"])],
         "rows": venue_rows},
        {"name": "Units", "purpose": f"One row per {v['group']}: the thing a {v['requirement']} belongs to.",
         "columns": [_column("Name", f"What the {v['group']} is called", units[0]),
                     _column("Size", f"How many people the {v['group']} covers, when that matters", CLASS_SIZE)],
         "rows": unit_rows},
        {"name": "Duties", "purpose": f"One row per {v['requirement']}: what has to be covered, how often "
                                       f"and by whom. The sample lists a few {v['requirement']}s; list every "
                                       f"{v['requirement']} your {v['group']} needs.",
         "columns": [_column("Name", f"What the {v['requirement']} is called", duties[0]["name"]),
                     _column("Unit", f"The {v['group']} it belongs to, from the Units sheet", units[0]),
                     _column("Per cycle", "How many times it happens in one cycle", per_cycle[0]),
                     _column("Length", f"How long it runs, in {_unit_word(template, knobs)}", duties[0]["length_slots"]),
                     _column("Venue kind", "The kind of place it needs, from the Venues sheet", duties[0]["venue_kind"]),
                     _column("Who can do it", f"The {v['person']}s who may take it, separated by commas, "
                                              f"or a role", f"{people[0]}, {people[1]}"),
                     _column("Together with", f"Another {v['requirement']} that must start at the same time", ""),
                     _column("Co-staffed", f"\"no\", or how many {v['person']}s take it together", duties[0]["min_staff"])],
         "rows": duty_rows},
    ]


def _deployment_sheets(template: dict, knobs: dict) -> list[dict]:
    """A sheet per subject family, three rows each: an entire-class row, then either the two options
    of a band or a second class and a co-taught row with a Split."""
    v = template["vocabulary"]
    sample = template["sample"]
    columns = _deployment_columns(v)
    head = [c["name"] for c in columns]
    levels = [(f"Year {n + 1}", [f"{n + 1}{_LETTERS[c]}" for c in range(sample["classes_per_level"])])
              for n in range(sample["levels"])]
    first_level, first_codes = levels[0]
    pairs = [(lvl, code) for lvl, codes in levels for code in codes]   # every (level, class) of the sample
    teachers = _people(v["person"], max(3, sample["teachers"]))
    options = _options_per_band(knobs)
    longest = _longest_run(rules_for(template, knobs))
    subject_codes = _subject_codes(sample["subjects"])

    families: dict[str, list[dict]] = {}
    for subject in sample["subjects"]:
        families.setdefault(str(subject["name"]).split()[0], []).append(subject)

    def row(level: str, subject: dict, grouping: str, classes: list[str], staff: list[str], split="") -> list:
        lengths = subject["lengths"]
        cells = [level, f"{subject['name']} ({subject['periods']} Periods)"]
        cells += [lengths.get(str(n)) or None for n in (1, 2, 3, 4)]
        cells += [subject["periods"], grouping]
        cells += [classes[i] if i < len(classes) else None for i in range(6)]
        cells += [CLASS_SIZE, staff[0], staff[1] if len(staff) > 1 else None, split or None]
        return cells

    # Teachers and (level, class) pairs are handed out round-robin across every sheet, not restarted
    # on each: the same first teacher on the first row of every family would carry the whole workbook
    # on their own, and the filled-in sample would not build.
    hand_out = _Roster(teachers)
    pair_at = 0

    def next_pair() -> tuple[str, str]:
        nonlocal pair_at
        pair = pairs[pair_at % len(pairs)]
        pair_at += 1
        return pair

    sheets, group_codes, banded_rows = [], [], set()
    for family, subjects in families.items():
        banded = next((s for s in subjects if s["band"]), None)
        first = subjects[0]
        level, code = next_pair()
        rows = [row(level, first, "Entire Class", [code], hand_out.team(1))]
        if banded is not None:
            codes_here = [f"1{subject_codes[banded['name']]}{_LETTERS[i]}" for i in range(max(2, options))]
            group_codes += codes_here[:2]
            for option_code in codes_here[:2]:
                banded_rows.add((family[:31], len(rows)))
                rows.append(row(first_level, banded, option_code, first_codes, hand_out.team(1)))
        else:
            other = subjects[-1]
            # a different class (and level, where the sample has one) each row, so no two rows are
            # the same requirement written twice
            second, third = next_pair(), next_pair()
            rows.append(row(second[0], other, "Entire Class", [second[1]], hand_out.team(1)))
            # a Split only makes sense where there is more than one lesson to share out, and only
            # where each teacher's share still fits a day
            split = (_split_of(other) if sum(other["lengths"].values()) > 1
                     and max(_split_parts(other["periods"])) <= longest else "")
            rows.append(row(third[0], other, "Entire Class", [third[1]], hand_out.team(2), split))
        sheets.append({"name": family[:31], "purpose": f"One row per {v['requirement']} taught to a "
                                                       f"{v['group']} or an option group.",
                       "columns": columns, "rows": rows, "head": head})

    if not any(r[-1] for s in sheets for r in s["rows"]):
        # every family is banded (a college of modules): put the Split on the first entire-class row
        # that has more than one lesson to share out, so the workbook still shows what the column does.
        # Never on an option row: splitting one makes two requirements for one option group, and the
        # band can no longer pair its options up.
        for sheet in sheets:
            for i, cells in enumerate(sheet["rows"]):
                if (sheet["name"], i) in banded_rows or cells[-2] is not None:
                    continue
                if sum(n or 0 for n in cells[2:6]) > 1 and max(_split_parts(cells[6])) <= longest:
                    cells[-2] = hand_out.team(1)[0]
                    cells[-1] = "%d/%d" % _split_parts(cells[6])
                    break
            else:
                continue
            break

    groups_columns = [_column(str(n), f"A {v['group']} of one division, one a column",
                              first_codes[n - 1] if n <= len(first_codes) else "") for n in range(1, 5)]
    groups_columns += [_column(f"Group {n}", "An option group the division splits into",
                               group_codes[n - 1] if n <= len(group_codes) else "") for n in (1, 2)]
    # only the first division has example option groups: the subject sheets name its codes
    groups_rows = [[*(cs + [None] * 4)[:4], *(group_codes[:2] if i == 0 else [])]
                   for i, (_lvl, cs) in enumerate(levels)]
    sheets.append({"name": "Groups", "purpose": "One row per division: the classes that split together, "
                                                "and the option groups they split into.",
                   "preamble": [["Class(es) in a Division", None, None, None, "Groups (Subjects)", None]],
                   "columns": groups_columns, "head": [1, 2, 3, 4, "Group 1", "Group 2"],
                   "rows": groups_rows})

    staff_rows = [[float(i + 1), name, list(families)[i % len(families)][:2].upper(), "Full time",
                   1.0 if i < len(teachers) - 1 else 0.5, _initials(name)]
                  for i, name in enumerate(teachers)]
    sheets.append({"name": "Control", "purpose": f"One row per {v['person']}: who they are and how much they teach.",
                   "preamble": [["STAFF"]],
                   "columns": [_column("S/N", "A running number", 1.0),
                               _column("Staff", f"The {v['person']}'s name, exactly as the subject sheets spell it", teachers[0]),
                               _column("Dept", "Which department they belong to", list(families)[0][:2].upper()),
                               _column("Profile", "Full time, part time, and so on", "Full time"),
                               _column("Teaching Load Factor", "1 for full time, 0.5 for half, and so on", 1.0),
                               _column("Short", "A short form for the printed timetable", _initials(teachers[0]))],
                   "head": ["S/N", "Staff", "Dept", "Profile", "Teaching Load Factor", "Short"],
                   "rows": staff_rows})
    return sheets


def sheets_for(template: dict, knobs: dict) -> list[dict]:
    """Every sheet of the workbook: its name, what it is for, its columns and its sample rows."""
    sheets = (_deployment_sheets if template["sheets"] == "deployment" else _generic_sheets)(template, knobs)
    for sheet in sheets:
        sheet.setdefault("head", [c["name"] for c in sheet["columns"]])
        sheet.setdefault("preamble", [])
    return sheets


def chosen_knobs(template: dict, knobs: dict) -> list[dict]:
    """The knobs as the user answered them, in the template's order, for the README and the guide."""
    out = []
    for knob in template["knobs"]:
        value = knobs.get(knob["key"], knob["default"])
        if isinstance(value, bool):
            value = "yes" if value else "no"
        out.append({"key": knob["key"], "label": knob["label"], "help": knob["help"], "value": value})
    return out


def workbook_bytes(template: dict, knobs: dict) -> bytes:
    """The workbook to fill in: the template's sheets with a few sample rows on each sheet (the Staff
    sheet lists whoever the sample duties need), and a README that explains every column in the
    template's own words."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = sheets_for(template, knobs)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["name"])
        for row in sheet["preamble"]:
            ws.append(row)
        ws.append(list(sheet["head"]))
        for row in sheet["rows"]:
            ws.append(list(row))

    readme = wb.create_sheet("README")
    readme.append([f"{template['name']} — the workbook to fill in"])
    readme.append([template["summary"]])
    readme.append([])
    readme.append(["What you chose", "", ""])
    for knob in chosen_knobs(template, knobs):
        readme.append([knob["label"], knob["value"], knob["help"]])
    readme.append([])
    readme.append(["Column", "What to put", "Example"])
    for sheet in sheets:
        readme.append([f"Sheet: {sheet['name']}", sheet["purpose"], ""])
        for column in sheet["columns"]:
            readme.append([str(column["name"]), column["what"], column["example"]])
    readme.append([])
    readme.append(["The sample rows on each sheet are examples: replace them with your own."])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def guide_html(templates_env, template: dict, knobs: dict) -> str:
    """One printable page: the configuration chosen, the knobs, each sheet and column, three steps."""
    v = template["vocabulary"]
    steps = [
        f"Fill in the workbook: replace the sample rows on each sheet with your own. The sample shows "
        f"a few rows; list every {v['requirement']} your organisation needs.",
        "Drop the filled workbook back into the chat. The app reads it, lists anything that does "
        "not add up, and turns it into a plan.",
        "Generate the draft timetable, then build it. Print or export when it looks right.",
    ]
    return templates_env.env.get_template("wizard/guide.html").render(
        template=template, vocabulary=template["vocabulary"], knobs=chosen_knobs(template, knobs),
        facts=facts(template, knobs), sheets=sheets_for(template, knobs), steps=steps,
        generated=_date.today().strftime("%-d %b %Y"))


# ---------------------------------------------------------------------------
# the sample PDF
# ---------------------------------------------------------------------------

def _sample_grids(org: dict, template: dict) -> list[G.Grid]:
    if template["sheets"] == "deployment":
        grids = G.all_grids(org, "classes")[:2] + G.all_grids(org, "teachers")[:2] + G.all_grids(org, "rooms")[:2]
    else:
        grids = G.all_grids(org, "teachers")[:3] + G.all_grids(org, "rooms")[:3]
    return grids[:MAX_SAMPLE_GRIDS]


def sample_pdf(engine, template: dict, knobs: dict, settings: dict, templates_env=None) -> bytes:
    """Build the sample organisation with the engine and print a few of its grids. The PDF needs no
    Jinja; `templates_env` is accepted so every producer of the wizard takes the same arguments."""
    org = sample_organisation(template, knobs, settings)
    result = engine.build(org)
    if result.get("unplaced") or result.get("clashes"):
        raise WizardError(
            f"{template['id']}: the sample does not build — "
            f"{len(result.get('unplaced') or [])} unplaced, {len(result.get('clashes') or [])} clashes")
    built = result["organisation"]
    return P.render(_sample_grids(built, template), _date.today().strftime("%-d %b %Y"), template["name"])
